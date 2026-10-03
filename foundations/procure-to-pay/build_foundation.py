"""Builds and verifies the procure-to-pay foundation document.

Run with:  ~/.cache/bwai-docx-venv/bin/python build_foundation.py
Written from general accounting and ERP process knowledge only; no company data used.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Pt, Cm

OUT = Path(__file__).with_name("foundation.docx")

HEADINGS = [
    "Purpose and how to use this document",
    "Business objects",
    "Quantities",
    "Tolerances",
    "Company parameters",
    "Rules",
    "What this document cannot decide",
]

# --------------------------------------------------------------------------- 1
PURPOSE = [
    "This document describes, before any file has been looked at, what must be true of a "
    "country affiliate's procure-to-pay data if the data is what it claims to be. It covers a "
    "narrow chain: purchase order, goods receipt, vendor invoice, payment, the clearing account "
    "that sits between goods receipt and invoice (goods received, not yet invoiced), the manual "
    "journals posted on that account, the mapping of local accounts to group accounts, and the "
    "exchange rates used to translate local currency into EUR. Three analyses depend on it: the "
    "balance of goods received but not yet invoiced at year end, invoices paid without a goods "
    "receipt, and the share of deliveries that arrived late.",
    "A software tool reads the tables below and tests each rule against the raw extracts in the "
    "group data warehouse. A rule that holds on almost all records is evidence that the extract "
    "is complete and that its columns mean what we think they mean. The records where a rule "
    "does not hold are not automatically errors: they are findings, listed for a person to look "
    "at. A rule that fails on a large share of records usually means the column was "
    "misunderstood or the extract is incomplete, not that the business is in disorder. No "
    "analysis should be trusted until the rules marked as applicable have been run.",
    "How to use it. As finance lead, read section 6 and write Yes or No in the column "
    "'Applies to us' for each rule; then supply the values in section 5. Leave everything else "
    "unchanged: the tables are read by a program, so the names in sections 2, 3 and 5 and the "
    "text in the column 'Formal' must not be reworded. If a rule does not fit your process, "
    "answer No rather than editing it, and say why in a covering note.",
    "Conventions used throughout. (a) Signs: quantities and amounts are signed. A goods receipt, "
    "an invoice and a payment are positive; a reversal, a return to vendor, a credit note and a "
    "returned payment are negative. Where the extract carries unsigned figures with a "
    "debit/credit or movement indicator, the sign is applied first (parameter "
    "DEBIT_CREDIT_REPRESENTATION). (b) On the clearing account a credit is positive and a debit "
    "negative, so that a positive total reads as goods received, not yet invoiced. (c) "
    "Quantities on goods receipts and invoices are in the order unit of the order line unless "
    "the name ends in _base. (d) Exchange rates are expressed as units of the target currency "
    "per one unit of the source currency, after applying RATE_QUOTATION. (e) Totals are compared "
    "per order line or per invoice over the whole life of the document up to CLOSING_DATE, "
    "which is why partial deliveries, several invoices per order line, reversals with "
    "re-posting, credit notes and later price changes do not break them. (f) A total is compared "
    "only for keys present on both sides; a key present on one side only is caught by the "
    "reference rules.",
]

# --------------------------------------------------------------------------- 2
OBJECTS = [
    ("purchase_order_line", "What was ordered from a vendor, at what price and for which date. Carries the current state of the line, including later price or quantity changes.",
     "One line of one purchase order (one material or service, one price)."),
    ("goods_receipt_line", "The physical receipt of ordered goods into the warehouse, and every reversal or return of such a receipt.",
     "One line of one goods movement document posted against an order line. A reversal or return is a separate record with negative quantity and value."),
    ("invoice_header", "The vendor's invoice or credit note as entered in invoice verification.",
     "One vendor invoice or credit note document. A reversal is a separate record with opposite sign."),
    ("invoice_line", "The part of a vendor invoice or credit note that belongs to one order line, or to no order line (for example freight).",
     "One line of one invoice or credit note."),
    ("payment", "The settlement of a vendor invoice by bank payment, offset or cash discount.",
     "One allocation of one payment document to one invoice. An invoice paid in two instalments has two records."),
    ("clearing_account_line", "Every ledger posting on the clearing account between goods receipt and invoice, whether made automatically by a goods receipt or invoice, or manually by journal.",
     "One line of one accounting document on a clearing account named in GRIR_ACCOUNTS."),
    ("grir_balance", "The open balance of goods received, not yet invoiced, as reported by the ERP at the closing date.",
     "One order line with its open quantity and open value at CLOSING_DATE."),
    ("vendor", "The supplier master: third parties and the group's supply company.",
     "One vendor account."),
    ("account_mapping", "The assignment of each local ledger account to an account of the group chart of accounts.",
     "One local account for one validity period."),
    ("exchange_rate", "The rates used to translate between currencies, including local currency to EUR.",
     "One rate for one currency pair, one rate type and one start date."),
]

# --------------------------------------------------------------------------- 3
QUANTITIES = [
    # purchase_order_line
    ("purchase_order_line", "po_line_id", "identifier", "Order number and line number together; the key every later document refers to."),
    ("purchase_order_line", "po_number", "identifier", "Number of the purchase order the line belongs to."),
    ("purchase_order_line", "vendor_id", "identifier", "Vendor the order was placed with."),
    ("purchase_order_line", "order_date", "date", "Date the order was created."),
    ("purchase_order_line", "requested_delivery_date", "date", "Date the goods were due; the yardstick for late delivery."),
    ("purchase_order_line", "ordered_quantity", "quantity", "Quantity ordered, in the order unit, as currently stored."),
    ("purchase_order_line", "units_per_order_unit", "rate", "Number of base units contained in one order unit (for example 10 packs per carton)."),
    ("purchase_order_line", "ordered_quantity_base", "quantity", "Quantity ordered expressed in the base unit of the material."),
    ("purchase_order_line", "net_price", "money", "Current net price in order currency for price_unit order units."),
    ("purchase_order_line", "price_unit", "quantity", "Number of order units the net price refers to (1, 100, 1000)."),
    ("purchase_order_line", "net_order_value", "money", "Net value of the line in order currency."),
    ("purchase_order_line", "order_currency", "text", "Currency of price and value."),
    ("purchase_order_line", "line_status", "status", "State of the line: open, delivery completed, deleted."),
    # goods_receipt_line
    ("goods_receipt_line", "gr_line_id", "identifier", "Movement document number, year and line together."),
    ("goods_receipt_line", "po_line_id", "identifier", "Order line the goods were received against."),
    ("goods_receipt_line", "movement_type", "status", "Kind of movement: receipt, reversal of a receipt, return to vendor."),
    ("goods_receipt_line", "posting_date", "date", "Date the movement was posted in the ledger; decides the period."),
    ("goods_receipt_line", "document_date", "date", "Date of the delivery note, that is when the goods physically arrived."),
    ("goods_receipt_line", "quantity", "quantity", "Signed quantity in the order unit: positive for a receipt, negative for a reversal or return."),
    ("goods_receipt_line", "quantity_base", "quantity", "The same quantity in the base unit, same sign."),
    ("goods_receipt_line", "value_local", "money", "Signed value in local currency posted to the clearing account by this movement."),
    ("goods_receipt_line", "reversed_gr_line_id", "identifier", "On a reversal: the receipt line it reverses. Empty otherwise."),
    ("goods_receipt_line", "ledger_document_id", "identifier", "Accounting document created by the movement. Empty for unvalued receipts."),
    # invoice_header
    ("invoice_header", "invoice_id", "identifier", "Internal number of the invoice document, including year."),
    ("invoice_header", "vendor_id", "identifier", "Vendor who issued the invoice."),
    ("invoice_header", "vendor_reference", "text", "The vendor's own invoice number."),
    ("invoice_header", "document_type", "status", "Invoice, credit note, or reversal."),
    ("invoice_header", "invoice_date", "date", "Date printed on the vendor's invoice."),
    ("invoice_header", "posting_date", "date", "Date the invoice was posted in the ledger."),
    ("invoice_header", "due_date", "date", "Date payment falls due under the payment terms."),
    ("invoice_header", "invoice_currency", "text", "Currency the invoice is issued in."),
    ("invoice_header", "net_amount", "money", "Signed net amount in invoice currency; negative for a credit note or reversal."),
    ("invoice_header", "tax_amount", "money", "Signed tax amount in invoice currency."),
    ("invoice_header", "gross_amount", "money", "Signed amount payable in invoice currency."),
    ("invoice_header", "exchange_rate", "rate", "Local currency units per one unit of invoice currency used at posting; 1 for local-currency invoices."),
    ("invoice_header", "gross_amount_local", "money", "Signed amount payable in local currency."),
    ("invoice_header", "settled_amount", "money", "Signed amount in invoice currency settled up to CLOSING_DATE by payment, offset or discount."),
    ("invoice_header", "reversed_invoice_id", "identifier", "On a reversal: the invoice it reverses. Empty otherwise."),
    # invoice_line
    ("invoice_line", "invoice_line_id", "identifier", "Invoice number and line number together."),
    ("invoice_line", "invoice_id", "identifier", "Invoice the line belongs to."),
    ("invoice_line", "po_line_id", "identifier", "Order line the invoice line refers to. Empty for lines without an order."),
    ("invoice_line", "quantity", "quantity", "Signed invoiced quantity in the order unit; negative on credit notes and reversals; zero for pure price adjustments."),
    ("invoice_line", "net_amount", "money", "Signed net amount in invoice currency."),
    ("invoice_line", "net_amount_local", "money", "Signed net amount in local currency."),
    ("invoice_line", "clearing_amount_local", "money", "Signed part of the line, in local currency, posted to the clearing account."),
    ("invoice_line", "variance_local", "money", "Signed part of the line, in local currency, posted as price or exchange-rate variance instead of to the clearing account."),
    # payment
    ("payment", "payment_id", "identifier", "Payment document number and allocation line together."),
    ("payment", "invoice_id", "identifier", "Invoice settled by this allocation."),
    ("payment", "payment_date", "date", "Date the payment was posted."),
    ("payment", "paid_amount", "money", "Signed amount paid in invoice currency."),
    ("payment", "cash_discount_amount", "money", "Signed cash discount taken, in invoice currency."),
    ("payment", "settled_amount", "money", "Signed amount of the invoice cleared by this allocation: paid plus discount."),
    ("payment", "paid_amount_local", "money", "Signed amount paid in local currency."),
    # clearing_account_line
    ("clearing_account_line", "ledger_line_id", "identifier", "Accounting document number, year and line together."),
    ("clearing_account_line", "ledger_document_id", "identifier", "Accounting document the line belongs to."),
    ("clearing_account_line", "local_account", "identifier", "Local ledger account posted to."),
    ("clearing_account_line", "document_type", "status", "Kind of accounting document: goods receipt, invoice, automatic clearing, manual journal."),
    ("clearing_account_line", "posting_date", "date", "Date that decides the accounting period."),
    ("clearing_account_line", "entry_date", "date", "Date the document was keyed in."),
    ("clearing_account_line", "po_line_id", "identifier", "Order line the posting belongs to."),
    ("clearing_account_line", "amount_local", "money", "Amount in local currency, credit positive and debit negative."),
    ("clearing_account_line", "amount_group", "money", "The same amount in EUR, same sign."),
    ("clearing_account_line", "line_text", "text", "Free-text explanation of the posting."),
    ("clearing_account_line", "entered_by", "identifier", "User or batch job that entered the document."),
    # grir_balance
    ("grir_balance", "po_line_id", "identifier", "Order line the open balance belongs to."),
    ("grir_balance", "closing_date", "date", "Date the balance is stated at."),
    ("grir_balance", "open_quantity", "quantity", "Quantity received minus quantity invoiced, in the order unit. Negative when invoiced ahead of receipt."),
    ("grir_balance", "balance_local", "money", "Open clearing-account balance in local currency; positive means received, not yet invoiced."),
    ("grir_balance", "closing_rate", "rate", "EUR per one unit of local currency at the closing date."),
    ("grir_balance", "balance_group", "money", "Open balance in EUR as reported to the group."),
    # vendor
    ("vendor", "vendor_id", "identifier", "Vendor account number."),
    ("vendor", "vendor_name", "text", "Name of the vendor."),
    ("vendor", "vendor_type", "status", "Whether the vendor is a group company or a third party."),
    ("vendor", "trading_partner", "identifier", "Group company code of the vendor, used for intercompany elimination. Empty for third parties."),
    # account_mapping
    ("account_mapping", "local_account", "identifier", "Account in the local chart of accounts."),
    ("account_mapping", "group_account", "identifier", "Account in the group chart of accounts it reports to."),
    ("account_mapping", "valid_from", "date", "First day the mapping applies."),
    ("account_mapping", "valid_to", "date", "Last day the mapping applies."),
    ("account_mapping", "mapping_status", "status", "Whether the mapping is active."),
    # exchange_rate
    ("exchange_rate", "from_currency", "text", "Currency being translated."),
    ("exchange_rate", "to_currency", "text", "Currency translated into."),
    ("exchange_rate", "rate_type", "status", "Closing rate, average rate, daily booking rate."),
    ("exchange_rate", "valid_from", "date", "First day the rate applies."),
    ("exchange_rate", "rate", "rate", "Units of to_currency per one unit of from_currency."),
]

# --------------------------------------------------------------------------- 4
TOL_PROSE = [
    "Figures that should agree rarely agree to the last digit. Amounts are rounded line by line "
    "and again in total; rates are stored with a limited number of decimals; a small difference "
    "on an invoice is often accepted and posted away. A tolerance says how much disagreement is "
    "still agreement. Each class has an absolute limit (in units of whatever is compared: "
    "currency units, units of quantity, days) and a relative limit (a fraction of the larger of "
    "the two figures). A comparison passes if the difference is within the absolute OR the "
    "relative limit: the absolute limit takes care of small figures, the relative limit of large "
    "ones. For a rule written with <= or >= the tolerance widens the permitted side only. "
    "Tolerances are deliberately tight: they absorb rounding, not business differences. The one "
    "exception is T-OVERRUN, which reflects the commercial over-delivery allowance and not rounding.",
]
TOLERANCES = [
    ("T-EXACT", "Identifiers, uniqueness, references, presence of a value, signs", "-", "-",
     "A key either matches or it does not; a sign is either right or wrong. Nothing to round."),
    ("T-MONEY", "Amounts in one currency compared with amounts in the same currency", "1.00", "0.0005",
     "Covers rounding per line against rounding in total and accepted small differences. One currency unit is immaterial in any currency; five hundredths of a percent keeps large documents honest."),
    ("T-FX", "Amounts compared after multiplying by an exchange rate", "1.00", "0.005",
     "Rates are stored with four or five significant decimals and amounts are rounded after conversion; half a percent absorbs this without hiding a wrong rate date or an inverted quotation."),
    ("T-QTY", "Quantities compared in the same unit or after unit conversion", "0.001", "0.001",
     "Quantities are held with three decimals; conversion between order unit and base unit rounds at the last decimal."),
    ("T-OVERRUN", "Received quantity against the ordered quantity", "-", "0.10",
     "Vendors may deliver slightly more than ordered within an agreed allowance; ten percent is the usual ceiling. Beyond that the order should have been changed."),
    ("T-DAYS", "Comparisons between two dates", "1", "-",
     "One day absorbs time-zone and end-of-day effects between the entry of a document and its dating."),
]
MIN_SHARE = "0.95"
MIN_SHARE_REASON = (
    "A rule counts as holding when at least this share of the records it was tested on satisfy "
    "it. The process has genuine exceptions - goods in transit invoiced before receipt, "
    "corrections by manual journal, lines closed by hand - and in a healthy affiliate they stay "
    "within a few percent. A misread column or an incomplete extract, by contrast, breaks a "
    "rule on a large share of records: a wrong sign convention fails about half, a missing "
    "table fails nearly all. Ninety-five percent separates the two cases cleanly. Rules that "
    "hold are still reported with their exceptions; rules below the threshold are reported as "
    "not holding and must be explained before any analysis is used."
)

# --------------------------------------------------------------------------- 5
PARAMETERS = [
    ("LOCAL_CURRENCY", "The currency the affiliate keeps its books in.", "CZK"),
    ("CLOSING_DATE", "The year-end date up to which documents are included and at which balances are stated.", "2025-12-31"),
    ("GRIR_ACCOUNTS", "The local account or accounts used as clearing account between goods receipt and invoice. Say whether intercompany and third-party purchases use separate accounts.", "191100; 191200"),
    ("MANUAL_JOURNAL_VALUE", "The value of document_type that marks a manually entered journal on the clearing account.", "SA"),
    ("REVERSAL_MOVEMENT_VALUE", "The value of movement_type that marks the reversal of a goods receipt.", "102"),
    ("INTERCOMPANY_VENDOR_VALUE", "The value of vendor_type that marks a group company.", "IC"),
    ("RATE_QUOTATION", "How exchange rates are quoted in the extract: local units per foreign unit (direct) or foreign units per local unit (indirect), and any ratio such as per 100.", "direct, per 1"),
    ("CLOSING_RATE_TYPE", "The value of rate_type that marks the group closing rate used at year end.", "CLOSING"),
    ("DEBIT_CREDIT_REPRESENTATION", "How debit and credit, and receipt and reversal, appear in the extract: by sign of the figure or by a separate indicator with unsigned figures.", "indicator S/H, amounts unsigned"),
]

# --------------------------------------------------------------------------- 6
GR_KEY = "goods_receipt_line.po_line_id"
RULES = [
    # (name, statement, kind, formal, tolerance, why, exception)
    ("Order line is unique",
     "Each order line appears once.",
     "UNIQUE", "purchase_order_line: po_line_id", "T-EXACT",
     "Every receipt, invoice and balance is totalled per order line. A duplicated line doubles every total built on it.",
     "The extract was joined to change history or schedule lines and multiplied the rows; or two company codes were mixed."),
    ("Goods receipt line is unique",
     "Each goods movement line appears once.",
     "UNIQUE", "goods_receipt_line: gr_line_id", "T-EXACT",
     "A receipt counted twice overstates goods received, not yet invoiced.",
     "Document year missing from the key; rows multiplied by a join to batches or serial numbers."),
    ("Invoice is unique",
     "Each invoice document appears once.",
     "UNIQUE", "invoice_header: invoice_id", "T-EXACT",
     "Payments and invoice lines are attached to the invoice. Duplicates would show invoices as only partly paid or over-invoiced.",
     "Fiscal year missing from the key; header joined to its lines or to tax lines."),
    ("One rate per currency pair, type and day",
     "For a currency pair, rate type and start date there is exactly one rate.",
     "UNIQUE", "exchange_rate: from_currency, to_currency, rate_type, valid_from", "T-EXACT",
     "Two rates for the same day make the translation into EUR arbitrary.",
     "A rate was corrected and both versions were extracted; rates from two source systems were combined."),
    ("One group account per local account and period",
     "A local account has one mapping per start date.",
     "UNIQUE", "account_mapping: local_account, valid_from", "T-EXACT",
     "A balance can be reported to the group once only. A double mapping reports it twice or splits it unpredictably.",
     "The mapping was changed during the year without closing the old entry; mappings for several reporting versions were extracted together."),
    ("Goods receipt belongs to an order line",
     "Every goods receipt refers to an order line that exists.",
     "REFERENCE", "goods_receipt_line.po_line_id -> purchase_order_line.po_line_id", "T-EXACT",
     "Value reaches the clearing account only through a receipt against an order. Without the order line there is no price, no due date and no vendor.",
     "Orders outside the extracted period or archived; receipts against orders of another company code; key built differently in the two extracts."),
    ("Invoice line belongs to an order line",
     "Every invoice line that names an order line names one that exists.",
     "REFERENCE", "invoice_line.po_line_id -> purchase_order_line.po_line_id", "T-EXACT",
     "The invoice clears the clearing account for a specific order line; the clearing cannot be followed if that line is unknown.",
     "Old orders missing from the extract; invoice entered against a wrong or deleted order line."),
    ("Invoiced order line has a goods receipt",
     "Every order line that has been invoiced has at least one goods receipt.",
     "REFERENCE", "invoice_line.po_line_id -> goods_receipt_line.po_line_id", "T-EXACT",
     "For goods the three-way match requires receipt before the invoice is released for payment. This rule is the basis for the analysis of invoices paid without goods receipt.",
     "Goods in transit, typical for intercompany supply invoiced at dispatch; service or cost lines not subject to goods receipt; a receipt that was never posted - the true finding."),
    ("Payment belongs to an invoice",
     "Every payment allocation refers to an invoice that exists.",
     "REFERENCE", "payment.invoice_id -> invoice_header.invoice_id", "T-EXACT",
     "A payment that cannot be traced to an invoice cannot be traced to an order or a receipt either.",
     "Down payments and payments on account; invoices posted directly in the ledger outside invoice verification; invoices of an earlier year not extracted."),
    ("Reversal points to a real receipt",
     "Every reversal names a goods receipt line that exists.",
     "REFERENCE", "goods_receipt_line.reversed_gr_line_id -> goods_receipt_line.gr_line_id", "T-EXACT",
     "A reversal undoes a specific receipt. If the original is missing, the net received quantity is understated.",
     "The original receipt lies before the extracted period; the reference was built without the document year."),
    ("Clearing account is mapped to the group",
     "Every local account posted on has a group account mapping.",
     "REFERENCE", "clearing_account_line.local_account -> account_mapping.local_account", "T-EXACT",
     "The year-end balance reaches the group balance sheet only through the mapping.",
     "A new local account was opened and not yet mapped; the mapping extract is from a different date than the ledger."),
    ("Order value is quantity times price",
     "The net value of an order line equals ordered quantity times net price, divided by the price unit.",
     "IDENTITY", "purchase_order_line: ordered_quantity * net_price / price_unit = net_order_value", "T-MONEY",
     "It confirms that quantity, price and price unit are understood correctly; all later value comparisons rest on them. It stays true after a price change because all three are the current values.",
     "Price unit not extracted or misread (price per 100 or 1000); scaled or conditional pricing; free-of-charge lines; value in a different currency than price."),
    ("Order unit converts to base unit",
     "Ordered quantity in the order unit times the conversion factor equals ordered quantity in the base unit.",
     "IDENTITY", "purchase_order_line: ordered_quantity * units_per_order_unit = ordered_quantity_base", "T-QTY",
     "Stock is kept in the base unit while orders and invoices are in the order unit. A wrong factor distorts every quantity comparison.",
     "Factor stored as a fraction of which only the numerator was extracted; factor changed on the material after the order was placed."),
    ("Receipt quantity and value have the same sign",
     "On a goods movement line, quantity and value are both positive or both negative.",
     "IDENTITY", "goods_receipt_line: quantity * value_local >= 0", "T-EXACT",
     "A receipt adds quantity and value; a reversal or return removes both. Opposite signs mean the sign was applied to one column only.",
     "Debit/credit indicator applied inconsistently in the extract; value-only corrections recorded as goods movements."),
    ("Invoice net plus tax is gross",
     "On each invoice or credit note, net amount plus tax equals the amount payable.",
     "IDENTITY", "invoice_header: net_amount + tax_amount = gross_amount", "T-MONEY",
     "Confirms that the three amounts are in one currency and carry the same sign, for credit notes too.",
     "Reverse-charge or import tax not part of the payable; withholding tax; unplanned delivery costs kept outside net."),
    ("Invoice translates at its rate",
     "The amount payable in invoice currency times the exchange rate equals the amount in local currency.",
     "IDENTITY", "invoice_header: gross_amount * exchange_rate = gross_amount_local", "T-FX",
     "Intercompany purchases are often invoiced in EUR or another foreign currency. This confirms how the rate is quoted and that local amounts are what the ledger holds.",
     "Rate quoted the other way round or per 100 units; local amount entered manually to match a customs document; rate fixed by a hedging contract."),
    ("Invoice is posted on or after its date",
     "The posting date of an invoice is not earlier than the date on the vendor's invoice.",
     "IDENTITY", "invoice_header: posting_date >= invoice_date", "T-DAYS",
     "An invoice cannot be booked before it is issued. Confirms the two dates are not swapped, which matters for cut-off at year end.",
     "Invoice back-dated into a closed period's successor to meet a deadline; date fields swapped; typing error in the year."),
    ("Invoice line splits into clearing and variance",
     "The local net amount of an invoice line equals the part posted to the clearing account plus the part posted as price or exchange-rate variance.",
     "IDENTITY", "invoice_line: clearing_amount_local + variance_local = net_amount_local", "T-MONEY",
     "The clearing account is relieved at the value the receipt put on it; anything the vendor charges above or below is variance. This is how prices changed after the order show up without leaving a balance.",
     "Part of the line posted to a further account (freight, small differences); clearing amount taken from a different document than the line."),
    ("Closing balance translates at the closing rate",
     "The open balance in local currency times the closing rate equals the balance reported in EUR.",
     "IDENTITY", "grir_balance: balance_local * closing_rate = balance_group", "T-FX",
     "The group reports in EUR. A liability is translated at the closing rate; this ties the local figure to the group figure.",
     "Balance translated at an average or historical rate; rate of a different date; group figure includes a top-side adjustment."),
    ("Reversal names what it reverses",
     "A goods movement marked as reversal carries the reference to the reversed receipt.",
     "CONDITION", "goods_receipt_line: when movement_type = REVERSAL_MOVEMENT_VALUE then reversed_gr_line_id present", "T-EXACT",
     "Without the link, a reversal and its re-posting cannot be told apart from a genuine return.",
     "REVERSAL_MOVEMENT_VALUE also covers returns to vendor; reference column not extracted."),
    ("Manual journal carries an explanation",
     "A manual journal on the clearing account has a line text.",
     "CONDITION", "clearing_account_line: when document_type = MANUAL_JOURNAL_VALUE then line_text present", "T-EXACT",
     "The account is meant to be moved by receipts and invoices only. A manual entry overrides that and, under internal control rules, must say why.",
     "Explanation held in the document header or an attachment only; recurring accrual posted by a job with a manual document type."),
    ("Manual journal names an order line",
     "A manual journal on the clearing account states the order line it corrects.",
     "CONDITION", "clearing_account_line: when document_type = MANUAL_JOURNAL_VALUE then po_line_id present", "T-EXACT",
     "The balance is explained order line by order line. An entry without an order line is a lump sum that no receipt or invoice will ever clear.",
     "Year-end lump-sum accrual or write-off; reclassification between clearing accounts; migration balance."),
    ("Intercompany vendor has a trading partner",
     "A vendor marked as group company carries the group company code.",
     "CONDITION", "vendor: when vendor_type = INTERCOMPANY_VENDOR_VALUE then trading_partner present", "T-EXACT",
     "Intercompany payables and goods in transit are eliminated at group level by trading partner. Without it the supply company's balance cannot be separated from third parties.",
     "Vendor set up in the wrong account group; trading partner maintained only on the document."),
    ("Received at most what was ordered",
     "Per order line, total quantity received net of reversals and returns does not exceed the ordered quantity.",
     "TOTAL", "sum(goods_receipt_line.quantity) by goods_receipt_line.po_line_id <= purchase_order_line.ordered_quantity by purchase_order_line.po_line_id", "T-OVERRUN",
     "A receipt is posted against the open order quantity. Any number of partial deliveries is normal; the net total is bounded by the order plus the over-delivery allowance.",
     "Receipt quantity in base unit while the order is in order unit; order quantity reduced after delivery; receipt posted against the wrong line."),
    ("Reversals reduce, never add",
     "Per order line, receipts without reversals are at least the net total, that is reversals sum to zero or less.",
     "TOTAL", "sum(goods_receipt_line.quantity where goods_receipt_line.movement_type != REVERSAL_MOVEMENT_VALUE) by goods_receipt_line.po_line_id >= sum(goods_receipt_line.quantity) by " + GR_KEY, "T-EXACT",
     "Confirms that REVERSAL_MOVEMENT_VALUE and the sign convention are understood: a reversal takes quantity away.",
     "Reversal quantities extracted unsigned; the value given marks receipts instead of reversals."),
    ("Invoiced at most what was received",
     "Per order line, total quantity invoiced net of credit notes does not exceed total quantity received net of reversals.",
     "TOTAL", "sum(invoice_line.quantity) by invoice_line.po_line_id <= sum(goods_receipt_line.quantity) by " + GR_KEY, "T-QTY",
     "Invoice verification matches the invoice against goods received. Several invoices for one line are normal; their net total stays within the net receipts.",
     "Invoice ahead of receipt (goods in transit); receipt reversed after the invoice was posted; duplicate invoice; quantity on price-adjustment lines not set to zero."),
    ("Invoice lines add up to the header",
     "Per invoice, the net amounts of its lines add up to the net amount of the header.",
     "TOTAL", "sum(invoice_line.net_amount) by invoice_line.invoice_id = invoice_header.net_amount by invoice_header.invoice_id", "T-MONEY",
     "Confirms that all lines of each invoice were extracted, including lines without an order.",
     "Lines posted directly to a ledger account not in the extract; unplanned delivery costs; small difference accepted at header level."),
    ("Payments add up to the settled amount",
     "Per invoice, the settled amounts of its payment allocations add up to the settled amount on the invoice.",
     "TOTAL", "sum(payment.settled_amount) by payment.invoice_id = invoice_header.settled_amount by invoice_header.invoice_id", "T-MONEY",
     "Confirms that the payment extract is complete, so that 'paid' can be relied on when looking for invoices paid without receipt. Partial payments are covered because totals are compared.",
     "Invoice offset against a credit note that is not in the payment extract; payment reset and re-cleared; payments after CLOSING_DATE included on one side only."),
    ("Ledger lines add up to the reported balance",
     "Per order line, all postings on the clearing account add up to the open balance reported at the closing date.",
     "TOTAL", "sum(clearing_account_line.amount_local) by clearing_account_line.po_line_id = grir_balance.balance_local by grir_balance.po_line_id", "T-MONEY",
     "The reported balance of goods received, not yet invoiced, is nothing other than the sum of what was posted. This ties the year-end figure to the ledger.",
     "Postings after the closing date included; ledger extract starts later than the oldest open item; balance list produced on a different day."),
    ("Open quantity is received minus invoiced",
     "Per order line, quantity received minus quantity invoiced equals the open quantity reported at the closing date.",
     "DIFFERENCE", "sum(goods_receipt_line.quantity) by goods_receipt_line.po_line_id - sum(invoice_line.quantity) by invoice_line.po_line_id = sum(grir_balance.open_quantity) by grir_balance.po_line_id", "T-QTY",
     "This is the definition of goods received, not yet invoiced, in quantity. It holds whatever the number of partial deliveries, invoices, reversals and credit notes.",
     "Receipts or invoices missing from the extract; a line cleared by account maintenance without an invoice; documents dated after the closing date."),
    ("Clearing account moves only by receipt and invoice",
     "Per order line, value received minus value cleared by invoices equals the clearing-account postings other than manual journals.",
     "DIFFERENCE", "sum(goods_receipt_line.value_local) by goods_receipt_line.po_line_id - sum(invoice_line.clearing_amount_local) by invoice_line.po_line_id = sum(clearing_account_line.amount_local where clearing_account_line.document_type != MANUAL_JOURNAL_VALUE) by clearing_account_line.po_line_id", "T-MONEY",
     "Ties the purchasing documents to the ledger. Whatever remains unexplained on an order line is, by construction, the effect of manual journals - exactly what a reviewer needs to see.",
     "Automatic clearing runs or account maintenance booked under a non-manual document type; MANUAL_JOURNAL_VALUE incomplete; revaluation of foreign-currency items; sign convention of the ledger extract reversed."),
]

# --------------------------------------------------------------------------- 7
CANNOT_DECIDE = [
    "From what age an open balance of goods received, not yet invoiced, should be released to profit, and who approves it. The data shows the age; it does not show whether the vendor will still invoice.",
    "Whether goods invoiced by the supply company but still in transit at year end belong to the affiliate's stock and liability. This depends on the delivery terms agreed in the group, not on any posting.",
    "Which purchases are legitimately exempt from goods receipt (services, rent, utilities, subscriptions), so that an invoice paid without receipt is by design and not a control failure.",
    "What counts as late: measured against the requested date or the date the vendor confirmed, with how many days of grace, and whether a partial delivery on time counts as on time.",
    "Whether a manual journal on the clearing account was justified. The data shows who posted it, when and with what text; whether the reason is acceptable is a judgement.",
    "Which over-delivery and price deviations are acceptable. The tolerances here confirm the data; they do not replace the company's own approval limits.",
    "Whether the mapping of the clearing account to its group account is the right one, in particular whether intercompany and third-party balances must be reported separately. The data shows that a mapping exists, not that it is correct.",
    "Which exchange rate the group policy requires for each purpose (closing, average, transaction date). The data shows which rate was used.",
    "Whether a debit balance on an order line (invoiced ahead of receipt) should be presented as an asset or netted against the liability.",
]

RULE_COLS = ["ID", "Name", "Statement", "Kind", "Formal", "Tolerance", "Why it must hold",
             "An exception usually means", "Applies to us"]
TABLE_HEADERS = {
    "Business objects": ["Object", "Meaning", "One record is"],
    "Quantities": ["Object", "Quantity", "Type", "Meaning"],
    "Tolerances": ["Class", "Applies to", "Absolute", "Relative", "Reasoning"],
    "Company parameters": ["Parameter", "Meaning", "Example"],
    "Rules": RULE_COLS,
}
TYPES = {"money", "quantity", "rate", "date", "identifier", "status", "text"}


def add_table(doc, header, rows, size=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, header):
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = ""
            c.paragraphs[0].add_run(v).font.size = Pt(size)
    return t


def build():
    doc = Document()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, side, Cm(1.5))
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Foundation document: procure-to-pay and the goods received, not yet invoiced account", 0)

    doc.add_heading(HEADINGS[0], 1)
    for p in PURPOSE:
        doc.add_paragraph(p)

    doc.add_heading(HEADINGS[1], 1)
    add_table(doc, TABLE_HEADERS["Business objects"], OBJECTS)

    doc.add_heading(HEADINGS[2], 1)
    add_table(doc, TABLE_HEADERS["Quantities"], QUANTITIES)

    doc.add_heading(HEADINGS[3], 1)
    for p in TOL_PROSE:
        doc.add_paragraph(p)
    add_table(doc, TABLE_HEADERS["Tolerances"], TOLERANCES)
    p = doc.add_paragraph()
    p.add_run("Minimum share").bold = True
    doc.add_paragraph("MIN_SHARE = " + MIN_SHARE)
    doc.add_paragraph(MIN_SHARE_REASON)

    doc.add_heading(HEADINGS[4], 1)
    add_table(doc, TABLE_HEADERS["Company parameters"], PARAMETERS)

    doc.add_heading(HEADINGS[5], 1)
    rows = [("R%02d" % (i + 1),) + r + ("",) for i, r in enumerate(RULES)]
    add_table(doc, RULE_COLS, rows, size=8)

    doc.add_heading(HEADINGS[6], 1)
    for item in CANNOT_DECIDE:
        doc.add_paragraph(item, style="List Bullet")
    doc.save(OUT)


# ------------------------------------------------------------------ verification
NAME = r"[a-z][a-z0-9_]*"
PARAM = r"[A-Z][A-Z0-9_]*"
SUM = (r"sum\((?P<o{n}>" + NAME + r")\.(?P<q{n}>" + NAME + r")"
       r"(?: where (?P<fo{n}>" + NAME + r")\.(?P<fq{n}>" + NAME + r") (?:!=|=) (?P<fp{n}>" + PARAM + r"))?\)"
       r" by (?P<ko{n}>" + NAME + r")\.(?P<kq{n}>" + NAME + r")")


def verify():
    errs = []
    doc = Document(OUT)
    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    if h1 != HEADINGS:
        errs.append("headings: %r" % h1)

    def rows_of(t):
        return [[c.text for c in r.cells] for r in t.rows]

    tables = {}
    for t in doc.tables:
        data = rows_of(t)
        for name, hdr in TABLE_HEADERS.items():
            if data[0] == hdr:
                tables[name] = data[1:]
    for name in TABLE_HEADERS:
        if name not in tables:
            errs.append("missing table with exact header: " + name)
    if errs:
        return errs, {}
    if len(doc.tables) != 5:
        errs.append("unexpected number of tables")

    objects = [r[0] for r in tables["Business objects"]]
    qty = {}
    for o, q, ty, _ in tables["Quantities"]:
        if o not in objects:
            errs.append("quantity for unknown object " + o)
        if ty not in TYPES:
            errs.append("bad type " + ty)
        if not re.fullmatch(NAME, q) or not re.fullmatch(NAME, o):
            errs.append("bad name " + o + "." + q)
        if q in qty.setdefault(o, {}):
            errs.append("duplicate quantity " + o + "." + q)
        qty[o][q] = ty
    for o in objects:
        if o not in qty:
            errs.append("object without quantities " + o)
    classes = [r[0] for r in tables["Tolerances"]]
    if "T-EXACT" not in classes or len(set(classes)) != len(classes):
        errs.append("tolerance classes")
    for r in tables["Tolerances"]:
        for v in r[2:4]:
            if v != "-":
                float(v)
    params = [r[0] for r in tables["Company parameters"]]
    for p in params:
        if not re.fullmatch(PARAM, p):
            errs.append("bad parameter " + p)
    ms = [p.text for p in doc.paragraphs if re.fullmatch(r"MIN_SHARE = 0\.\d\d", p.text)]
    if len(ms) != 1:
        errs.append("MIN_SHARE line")

    def has(o, q, rid, ty=None):
        if o not in qty or q not in qty[o]:
            errs.append("%s: undefined %s.%s" % (rid, o, q))
        elif ty and qty[o][q] != ty:
            errs.append("%s: %s.%s is not %s" % (rid, o, q, ty))

    def par(p, rid):
        if p not in params:
            errs.append("%s: undefined parameter %s" % (rid, p))

    def sumterm(m, n, rid):
        o = m.group("o%d" % n)
        has(o, m.group("q%d" % n), rid)
        if m.group("ko%d" % n) != o:
            errs.append(rid + ": key object differs from summed object")
        has(o, m.group("kq%d" % n), rid)
        if m.group("fo%d" % n):
            if m.group("fo%d" % n) != o:
                errs.append(rid + ": filter object differs")
            has(o, m.group("fq%d" % n), rid, "status")
            par(m.group("fp%d" % n), rid)

    kinds = {}
    for i, r in enumerate(tables["Rules"]):
        rid, name, stmt, kind, formal, tol, why, exc, applies = r
        if rid != "R%02d" % (i + 1):
            errs.append("bad id " + rid)
        if applies != "":
            errs.append(rid + ": Applies to us not empty")
        if tol not in classes:
            errs.append(rid + ": tolerance " + tol)
        if "\n" in formal or not all([name, stmt, why, exc]):
            errs.append(rid + ": empty cell or multi-line formal")
        kinds[kind] = kinds.get(kind, 0) + 1
        if kind == "IDENTITY":
            m = re.fullmatch(r"(" + NAME + r"): (.+?) (<=|>=|=) (.+)", formal)
            if not m:
                errs.append(rid + ": syntax"); continue
            o = m.group(1)
            for side in (m.group(2), m.group(4)):
                if re.search(r"[<>=]", side):
                    errs.append(rid + ": more than one operator")
                toks = re.findall(NAME + r"|\d+(?:\.\d+)?|[-+*/()]|\S+", side)
                for tk in toks:
                    if re.fullmatch(NAME, tk):
                        has(o, tk, rid)
                    elif not re.fullmatch(r"\d+(?:\.\d+)?|[-+*/()]", tk):
                        errs.append("%s: bad token %s" % (rid, tk))
                try:
                    compile(re.sub(NAME, "1", side), "x", "eval")
                except SyntaxError:
                    errs.append(rid + ": expression does not parse")
        elif kind == "UNIQUE":
            m = re.fullmatch(r"(" + NAME + r"): (" + NAME + r"(?:, " + NAME + r")*)", formal)
            if not m:
                errs.append(rid + ": syntax"); continue
            for q in m.group(2).split(", "):
                has(m.group(1), q, rid)
        elif kind == "REFERENCE":
            m = re.fullmatch(r"(%s)\.(%s) -> (%s)\.(%s)" % ((NAME,) * 4), formal)
            if not m:
                errs.append(rid + ": syntax"); continue
            has(m.group(1), m.group(2), rid)
            has(m.group(3), m.group(4), rid)
        elif kind == "CONDITION":
            m = re.fullmatch(r"(%s): when (%s) = (%s) then (%s) (present|absent)" % (NAME, NAME, PARAM, NAME), formal)
            if not m:
                errs.append(rid + ": syntax"); continue
            has(m.group(1), m.group(2), rid, "status")
            par(m.group(3), rid)
            has(m.group(1), m.group(4), rid)
        elif kind == "TOTAL":
            a = re.fullmatch(SUM.format(n=1) + r" (?:<=|>=|=) (?P<so>%s)\.(?P<sq>%s) by (?P<sko>%s)\.(?P<skq>%s)" % ((NAME,) * 4), formal)
            b = re.fullmatch(SUM.format(n=1) + r" (?:<=|>=|=) " + SUM.format(n=2), formal)
            if b:
                sumterm(b, 1, rid); sumterm(b, 2, rid)
            elif a:
                sumterm(a, 1, rid)
                if a.group("so") != a.group("sko"):
                    errs.append(rid + ": stored figure and key from different objects")
                has(a.group("so"), a.group("sq"), rid)
                has(a.group("so"), a.group("skq"), rid)
            else:
                errs.append(rid + ": syntax")
        elif kind == "DIFFERENCE":
            m = re.fullmatch(SUM.format(n=1) + " - " + SUM.format(n=2) + " = " + SUM.format(n=3), formal)
            if not m:
                errs.append(rid + ": syntax"); continue
            for n in (1, 2, 3):
                sumterm(m, n, rid)
        else:
            errs.append(rid + ": unknown kind " + kind)
    stats = dict(objects=len(objects), quantities=len(tables["Quantities"]), classes=len(classes),
                 parameters=len(params), rules=len(tables["Rules"]), kinds=kinds, min_share=ms)
    return errs, stats


if __name__ == "__main__":
    build()
    errors, stats = verify()
    print(stats)
    if errors:
        print("FAILED")
        for e in errors:
            print(" -", e)
        sys.exit(1)
    print("VERIFIED", OUT)
