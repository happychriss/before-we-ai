#!/usr/bin/env python3
"""Generate and verify the shipbuilding foundation document.

Run with:  ~/.cache/bwai-docx-venv/bin/python build_foundation.py
Writes foundation.docx next to this script, then re-opens and verifies it.
Content is general industry / accounting knowledge only; no company data.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path(__file__).resolve().parent / "foundation.docx"

HEADINGS = [
    "Purpose and how to use this document",
    "Business objects",
    "Quantities",
    "Tolerances",
    "Company parameters",
    "Rules",
    "What this document cannot decide",
]

H_OBJECTS = ["Object", "Meaning", "One record is"]
H_QUANT = ["Object", "Quantity", "Type", "Meaning"]
H_TOL = ["Class", "Applies to", "Absolute", "Relative", "Reasoning"]
H_PARAM = ["Parameter", "Meaning", "Example"]
H_RULES = ["ID", "Name", "Statement", "Kind", "Formal", "Tolerance",
           "Why it must hold", "An exception usually means", "Applies to us"]

TYPES = {"money", "quantity", "rate", "date", "identifier", "status", "text"}
KINDS = {"IDENTITY", "UNIQUE", "REFERENCE", "CONDITION", "TOTAL"}

# --------------------------------------------------------------------------
# 1. Purpose
# --------------------------------------------------------------------------
PURPOSE = [
    "This document describes what must be true of the accounting and operational "
    "data of a medium-sized shipbuilder that builds vessels to order, with a parent "
    "company (headquarters and yard) in one country and a wholly owned subsidiary "
    "yard in another country with a different currency. It was written before any "
    "of the company's files were seen and rests only on general shipbuilding and "
    "accounting practice. Its job is to be a yardstick: before anyone lets an AI "
    "answer questions such as \"What is the cost of building a vessel?\", \"What "
    "were our sales last year?\" or \"What is revenue per ship?\", a checking tool "
    "compares the raw ledgers, spreadsheets and exports against the rules below.",

    "A rule that holds on almost all records confirms that a file is what it claims "
    "to be and that its columns mean what we think they mean. An identity that ties "
    "four columns together and holds on 99 % of rows tells us something about all "
    "four columns at once. The records where a rule does not hold are not errors by "
    "definition; they are findings for a person to look at. A rule that fails on "
    "most records usually means the file was misread (wrong column, wrong sign, "
    "wrong currency, wrong level of detail), not that the business is wrong.",

    "How to use it. As finance lead, read the Business objects and Quantities "
    "sections and check that each one exists somewhere in your systems, whatever it "
    "is called locally. Fill in the Company parameters. Then go through the Rules "
    "table and write Yes, No or Partly in the column \"Applies to us\"; where you "
    "write No, a short reason is useful. Finally answer the policy questions in the "
    "last section: no data check can settle them, yet every one of the three "
    "business questions depends on them.",

    "Conventions. Object, quantity and parameter names are fixed vocabulary and are "
    "read by a program, as are the columns Kind, Formal and Tolerance; please do "
    "not reword them. Money is signed: credit notes, reversals and cost corrections "
    "carry a negative sign. Shares and percentages are fractions between 0 and 1. "
    "An exchange rate is the number of units of the target (functional or buyer) "
    "currency for one unit of the source currency. An empty cell means \"no "
    "value\", which is different from zero. A rule is only evaluated on records "
    "where all quantities it needs are present, except for CONDITION rules, which "
    "test presence itself.",
]

# --------------------------------------------------------------------------
# 2. Business objects
# --------------------------------------------------------------------------
OBJECTS = [
    ("legal_entity", "A company of the group that keeps its own books: the parent and the subsidiary yard.",
     "one legal entity with its own functional currency"),
    ("customer", "The shipowner or other party that orders a vessel.", "one customer"),
    ("vessel_project", "A newbuilding: one vessel built to order, known in the yard by its hull number. The cost object on which everything is collected.",
     "one vessel (hull) being built or already built"),
    ("shipbuilding_contract", "The sales contract under which a vessel is built and handed over, with its price as amended.",
     "one contract for one vessel"),
    ("change_order", "An amendment to a contract that changes the price (owner's extras, deletions, agreed variations).",
     "one priced amendment to one contract"),
    ("payment_milestone", "The contractual instalment plan: the events (signing, steel cutting, keel laying, launch, delivery) at which a part of the price falls due.",
     "one instalment of one contract"),
    ("sales_invoice", "An invoice or credit note issued to a customer, in shipbuilding normally one per instalment.",
     "one invoice or credit note"),
    ("cost_ledger_entry", "The project cost ledger (job cost ledger): every amount charged to a vessel, whatever its source.",
     "one cost posting to one vessel"),
    ("project_cost_summary", "The project controller's status report per vessel: cost to date by category, hours, estimate to complete, stage of completion and revenue recognised.",
     "one vessel at the reporting date"),
    ("employee", "Master list of the yard's own workers and staff who book time.", "one employee"),
    ("timesheet_entry", "Hours booked by an employee to a vessel and work package, valued at a costing rate.",
     "one booking of hours by one employee on one day to one vessel"),
    ("purchase_order_line", "An ordered item or service from a supplier, in shipbuilding mostly ordered directly for a specific vessel (main engine, steel, outfitting, subcontract work).",
     "one line of one purchase order"),
    ("material_issue", "A withdrawal of stock material (plate, profiles, pipe, consumables) from the store to a vessel.",
     "one issue of one item to one vessel"),
    ("account", "The chart of accounts.", "one general ledger account"),
    ("journal_voucher", "The header of a general ledger posting document.", "one posting document of one entity"),
    ("general_ledger_entry", "The general ledger: the debit and credit lines of each posting document.",
     "one line of one posting document"),
    ("intercompany_transaction", "A delivery or service between the two yards (for example hull blocks built by the subsidiary for a parent vessel), recorded in both currencies.",
     "one charge from one entity to the other"),
]

# --------------------------------------------------------------------------
# 3. Quantities
# --------------------------------------------------------------------------
QUANTITIES = [
    ("legal_entity", "entity_id", "identifier", "Key of the legal entity."),
    ("legal_entity", "entity_name", "text", "Name of the legal entity."),
    ("legal_entity", "country", "text", "Country in which the entity and its yard are located."),
    ("legal_entity", "functional_currency", "identifier", "Currency code in which the entity keeps its books."),

    ("customer", "customer_id", "identifier", "Key of the customer."),
    ("customer", "customer_name", "text", "Name of the customer."),
    ("customer", "customer_country", "text", "Country of the customer."),

    ("vessel_project", "project_id", "identifier", "Key of the vessel project as used on cost postings."),
    ("vessel_project", "hull_number", "identifier", "Yard (hull / newbuilding) number of the vessel."),
    ("vessel_project", "entity_id", "identifier", "Legal entity that builds the vessel and carries its cost."),
    ("vessel_project", "vessel_type", "text", "Type or class of vessel."),
    ("vessel_project", "project_status", "status", "Stage of the project (for example planned, in build, delivered)."),
    ("vessel_project", "keel_laying_date", "date", "Date of keel laying / start of hull assembly."),
    ("vessel_project", "launch_date", "date", "Date of launch or float-out."),
    ("vessel_project", "planned_delivery_date", "date", "Delivery date currently planned."),
    ("vessel_project", "delivery_date", "date", "Actual date of handover to the owner; empty until handover."),

    ("shipbuilding_contract", "contract_id", "identifier", "Key of the contract."),
    ("shipbuilding_contract", "project_id", "identifier", "Vessel that the contract is for."),
    ("shipbuilding_contract", "customer_id", "identifier", "Customer who ordered the vessel."),
    ("shipbuilding_contract", "contract_date", "date", "Date the contract was signed."),
    ("shipbuilding_contract", "contract_currency", "identifier", "Currency code in which the price is agreed."),
    ("shipbuilding_contract", "contract_price_original", "money", "Price at signing, net of tax, in contract currency."),
    ("shipbuilding_contract", "change_order_total", "money", "Net sum of all price changes to date, in contract currency."),
    ("shipbuilding_contract", "contract_price_current", "money", "Price as amended, net of tax, in contract currency."),
    ("shipbuilding_contract", "invoiced_total", "money", "Net amount invoiced to date under the contract, in contract currency."),

    ("change_order", "change_order_id", "identifier", "Key of the change order."),
    ("change_order", "contract_id", "identifier", "Contract that is amended."),
    ("change_order", "change_order_status", "status", "State of the change order (for example requested, approved, rejected)."),
    ("change_order", "approval_date", "date", "Date the owner agreed the change; empty until agreed."),
    ("change_order", "change_amount", "money", "Price change, positive for extras and negative for deletions, in contract currency."),
    ("change_order", "change_description", "text", "What is changed."),

    ("payment_milestone", "milestone_id", "identifier", "Key of the instalment."),
    ("payment_milestone", "contract_id", "identifier", "Contract the instalment belongs to."),
    ("payment_milestone", "milestone_sequence", "quantity", "Running number of the instalment within the contract."),
    ("payment_milestone", "milestone_event", "text", "Event that makes the instalment due (signing, steel cutting, keel laying, launch, delivery)."),
    ("payment_milestone", "milestone_amount", "money", "Instalment amount, net of tax, in contract currency."),
    ("payment_milestone", "due_date", "date", "Date the instalment falls or fell due."),

    ("sales_invoice", "invoice_id", "identifier", "Technical key of the invoice."),
    ("sales_invoice", "entity_id", "identifier", "Legal entity that issued the invoice."),
    ("sales_invoice", "invoice_number", "identifier", "Invoice number printed on the document."),
    ("sales_invoice", "contract_id", "identifier", "Contract the invoice is issued under; empty for sales outside a newbuilding contract."),
    ("sales_invoice", "invoice_date", "date", "Date of the invoice."),
    ("sales_invoice", "invoice_currency", "identifier", "Currency code of the invoice."),
    ("sales_invoice", "net_amount", "money", "Amount before tax in invoice currency; negative on credit notes."),
    ("sales_invoice", "tax_amount", "money", "VAT or sales tax in invoice currency; often zero for export deliveries."),
    ("sales_invoice", "gross_amount", "money", "Amount payable in invoice currency."),
    ("sales_invoice", "fx_rate", "rate", "Units of the issuing entity's functional currency per one unit of invoice currency; 1 if they are the same."),
    ("sales_invoice", "net_amount_functional", "money", "Amount before tax in the issuing entity's functional currency."),

    ("cost_ledger_entry", "cost_entry_id", "identifier", "Key of the cost posting."),
    ("cost_ledger_entry", "entity_id", "identifier", "Legal entity in whose books the cost is posted."),
    ("cost_ledger_entry", "project_id", "identifier", "Vessel that is charged."),
    ("cost_ledger_entry", "wbs_code", "text", "Work package or section of the build (hull, outfitting, machinery, design ...)."),
    ("cost_ledger_entry", "cost_category", "status", "Kind of cost (labour, material, subcontract, overhead, other)."),
    ("cost_ledger_entry", "posting_date", "date", "Date the cost was posted."),
    ("cost_ledger_entry", "transaction_currency", "identifier", "Currency code of the underlying document."),
    ("cost_ledger_entry", "amount_transaction", "money", "Cost in the currency of the underlying document."),
    ("cost_ledger_entry", "fx_rate", "rate", "Units of functional currency per one unit of transaction currency; 1 if they are the same."),
    ("cost_ledger_entry", "amount_functional", "money", "Cost in the posting entity's functional currency."),
    ("cost_ledger_entry", "source_document", "text", "Reference to the supplier invoice, timesheet, store issue or journal behind the posting."),

    ("project_cost_summary", "project_id", "identifier", "Vessel the summary is for."),
    ("project_cost_summary", "labour_hours_total", "quantity", "Own hours booked to the vessel to date."),
    ("project_cost_summary", "labour_cost_total", "money", "Own labour cost to date, in the building entity's functional currency."),
    ("project_cost_summary", "material_cost_total", "money", "Material and equipment cost to date."),
    ("project_cost_summary", "subcontract_cost_total", "money", "Subcontractor and bought-in service cost to date."),
    ("project_cost_summary", "overhead_cost_total", "money", "Overhead allocated to the vessel to date."),
    ("project_cost_summary", "other_cost_total", "money", "All other cost charged to the vessel to date."),
    ("project_cost_summary", "total_cost", "money", "Total cost charged to the vessel to date."),
    ("project_cost_summary", "estimated_cost_to_complete", "money", "Controller's estimate of cost still to come; zero after completion."),
    ("project_cost_summary", "estimated_cost_at_completion", "money", "Forecast of total cost of the vessel when finished."),
    ("project_cost_summary", "percent_complete", "rate", "Stage of completion as a fraction between 0 and 1."),
    ("project_cost_summary", "contract_revenue_functional", "money", "Current contract price expressed in the building entity's functional currency."),
    ("project_cost_summary", "revenue_recognised_total", "money", "Revenue recognised on the vessel to date, in functional currency."),

    ("employee", "employee_id", "identifier", "Key of the employee."),
    ("employee", "entity_id", "identifier", "Legal entity that employs the person."),
    ("employee", "cost_centre", "text", "Department or workshop the employee belongs to."),
    ("employee", "trade", "text", "Trade or function (welder, fitter, electrician, engineer ...)."),

    ("timesheet_entry", "timesheet_id", "identifier", "Key of the time booking."),
    ("timesheet_entry", "employee_id", "identifier", "Employee who worked."),
    ("timesheet_entry", "project_id", "identifier", "Vessel worked on; empty for non-project time."),
    ("timesheet_entry", "wbs_code", "text", "Work package worked on."),
    ("timesheet_entry", "work_date", "date", "Day the work was done."),
    ("timesheet_entry", "regular_hours", "quantity", "Hours at normal time."),
    ("timesheet_entry", "overtime_hours", "quantity", "Hours at overtime."),
    ("timesheet_entry", "total_hours", "quantity", "All hours of the booking."),
    ("timesheet_entry", "hourly_rate", "rate", "Costing rate per hour in the employing entity's functional currency."),
    ("timesheet_entry", "labour_cost", "money", "Value of the booking charged to the vessel."),

    ("purchase_order_line", "po_id", "identifier", "Purchase order number."),
    ("purchase_order_line", "po_line_number", "identifier", "Line number within the purchase order."),
    ("purchase_order_line", "supplier_id", "identifier", "Supplier the order is placed with."),
    ("purchase_order_line", "project_id", "identifier", "Vessel the item is ordered for; empty for stock or general purchases."),
    ("purchase_order_line", "order_date", "date", "Date of the order."),
    ("purchase_order_line", "po_currency", "identifier", "Currency code of the order."),
    ("purchase_order_line", "ordered_quantity", "quantity", "Quantity ordered."),
    ("purchase_order_line", "unit_price", "rate", "Net price per unit in order currency."),
    ("purchase_order_line", "line_amount", "money", "Net value of the line in order currency."),

    ("material_issue", "issue_id", "identifier", "Key of the store issue."),
    ("material_issue", "project_id", "identifier", "Vessel the material is issued to."),
    ("material_issue", "item_id", "identifier", "Stock item issued."),
    ("material_issue", "issue_date", "date", "Date of the issue."),
    ("material_issue", "issued_quantity", "quantity", "Quantity issued; negative for returns to store."),
    ("material_issue", "unit_cost", "rate", "Stock valuation price per unit at the time of issue."),
    ("material_issue", "issue_value", "money", "Value charged to the vessel."),

    ("account", "account_id", "identifier", "Account number."),
    ("account", "account_name", "text", "Account name."),
    ("account", "account_type", "status", "Class of account (asset, liability, equity, revenue, expense)."),

    ("journal_voucher", "journal_id", "identifier", "Key of the posting document."),
    ("journal_voucher", "entity_id", "identifier", "Legal entity in whose books the document is posted."),
    ("journal_voucher", "posting_date", "date", "Posting date, which determines the accounting period."),
    ("journal_voucher", "total_debit", "money", "Sum of the debit lines in functional currency."),
    ("journal_voucher", "total_credit", "money", "Sum of the credit lines in functional currency."),

    ("general_ledger_entry", "journal_id", "identifier", "Posting document the line belongs to."),
    ("general_ledger_entry", "line_number", "identifier", "Line number within the posting document."),
    ("general_ledger_entry", "account_id", "identifier", "Account posted to."),
    ("general_ledger_entry", "project_id", "identifier", "Vessel the line relates to, if any."),
    ("general_ledger_entry", "debit_amount", "money", "Debit amount in functional currency; zero or empty on credit lines."),
    ("general_ledger_entry", "credit_amount", "money", "Credit amount in functional currency; zero or empty on debit lines."),
    ("general_ledger_entry", "entry_text", "text", "Posting text."),

    ("intercompany_transaction", "ic_id", "identifier", "Key of the intercompany charge."),
    ("intercompany_transaction", "selling_entity_id", "identifier", "Entity that delivers and charges."),
    ("intercompany_transaction", "buying_entity_id", "identifier", "Entity that receives and is charged."),
    ("intercompany_transaction", "project_id", "identifier", "Vessel the delivery is for."),
    ("intercompany_transaction", "transaction_date", "date", "Date of the charge."),
    ("intercompany_transaction", "amount_seller", "money", "Amount in the selling entity's functional currency."),
    ("intercompany_transaction", "fx_rate", "rate", "Units of the buyer's currency per one unit of the seller's currency."),
    ("intercompany_transaction", "amount_buyer", "money", "Amount in the buying entity's functional currency."),
]

# --------------------------------------------------------------------------
# 4. Tolerances
# --------------------------------------------------------------------------
TOL_PROSE = [
    "Figures that describe the same thing do not always agree to the last cent, and "
    "a check that demands this will drown the reader in findings that mean nothing. "
    "Three things cause harmless differences. First, rounding: every document rounds "
    "tax, line values and converted amounts to two decimals, and a total built from "
    "a thousand rounded lines can legitimately differ by a few currency units from "
    "a total rounded once. Second, exchange rates: rates are stored with four to six "
    "decimals, and a booking rate may be a daily, monthly or contract rate, so a "
    "recomputed conversion differs slightly from the stored one. Third, estimates: "
    "stage of completion and cost to complete are judgements updated at period end, "
    "so figures derived from them are only as sharp as the estimate.",

    "Keys, document numbers, statuses and the presence of a value have no such "
    "excuse and must match exactly. The same is true of the balance of a posting "
    "document. The principle is therefore: the tolerance belongs to the kind of "
    "comparison, not to the individual rule. A few classes are defined below and "
    "every rule uses exactly one. Absolute tolerances are in the unit of the "
    "quantity compared (currency units of the record's own currency, or hours); "
    "relative tolerances are fractions of the larger of the two figures. A "
    "comparison passes if it is within the absolute OR the relative tolerance.",
]

TOLERANCES = [
    ("T-EXACT", "Keys, references, uniqueness, presence or absence of a value, and the debit/credit balance of a posting document.",
     "-", "-", "Nothing is rounded or estimated here; any difference is a real difference."),
    ("T-MONEY", "Arithmetic within a single record in a single currency (net plus tax, quantity times price, price plus changes).",
     "0.05", "-", "Each component is rounded to two decimals; a handful of rounding steps can add up to a few cents but never more."),
    ("T-FX", "A stored amount compared with the same amount recomputed through an exchange rate.",
     "1.00", "0.01", "Rates carry limited decimals and the booking rate may be a daily, monthly average or agreed rate; differences up to about one per cent are rate effects, larger ones are a wrong currency or an inverted rate."),
    ("T-HOURS", "Arithmetic on hours within a single time booking.",
     "0.25", "-", "Time is commonly booked in quarter hours or in decimal hours converted from minutes."),
    ("T-TOTAL", "A sum over many records compared with a stored total (money or hours).",
     "1.00", "0.001", "Accumulated rounding over many lines and a few records posted around the cut-off date; one tenth of one per cent is immaterial for the questions asked and still catches a missing or duplicated batch."),
    ("T-ESTIMATE", "Figures that depend on a forecast: stage of completion, cost at completion, revenue recognised over time.",
     "-", "0.02", "The inputs are judgements revised at each period end and the stored percentage is usually rounded to whole or tenth per cents."),
]

MIN_SHARE = "0.95"
MIN_SHARE_PROSE = (
    "A rule counts as holding when at least the share of evaluated records stated "
    "below satisfies it. Real ledgers always contain some corrections, reversals, "
    "migrated legacy records and manual journals, so demanding 100 % would reject "
    "every genuine data set. On the other hand a rule that fails on more than about "
    "one record in twenty is no longer describing the data: either the column means "
    "something else or the process is different. The remaining records, at most "
    "five per cent, are few enough for a person to review one by one. T-EXACT rules "
    "on keys should in practice come out far above this threshold; anything below "
    "0.99 there deserves a look even though the rule formally holds."
)

# --------------------------------------------------------------------------
# 5. Company parameters
# --------------------------------------------------------------------------
PARAMETERS = [
    ("STATUS_DELIVERED", "Value of project_status that means the vessel has been handed over to the owner.", "Delivered"),
    ("STATUS_IN_BUILD", "Value of project_status that means the vessel is under construction and not yet handed over.", "In production"),
    ("STATUS_CHANGE_APPROVED", "Value of change_order_status that means the owner has agreed the price change.", "Approved"),
    ("PARENT_ENTITY_ID", "Value of entity_id for the parent company.", "1000"),
    ("SUBSIDIARY_ENTITY_ID", "Value of entity_id for the subsidiary yard.", "2000"),
    ("PARENT_CURRENCY", "Functional currency of the parent.", "EUR"),
    ("SUBSIDIARY_CURRENCY", "Functional currency of the subsidiary.", "PLN"),
    ("GROUP_CURRENCY", "Currency in which group figures are to be reported.", "EUR"),
    ("FX_RATE_DIRECTION", "How exchange rates are stored: DIRECT (functional units per one foreign unit, as assumed here) or INVERSE.", "DIRECT"),
    ("FISCAL_YEAR_START_MONTH", "First month of the financial year, needed to decide what \"last year\" means.", "1"),
    ("PERCENT_SCALE", "Whether stage of completion is stored as a fraction (1) or in per cent (100).", "1"),
]

# --------------------------------------------------------------------------
# 6. Rules  (name, statement, kind, formal, tolerance, why, exception)
# --------------------------------------------------------------------------
RULES = [
    # --- contract and sales -------------------------------------------------
    ("Current contract price",
     "The current contract price is the original price plus the net total of change orders.",
     "IDENTITY",
     "shipbuilding_contract: contract_price_current = contract_price_original + change_order_total",
     "T-MONEY",
     "A shipbuilding price changes only through agreed amendments; revenue per ship is built on the current price.",
     "A price adjustment made outside the change-order process (currency or escalation clause, liquidated damages, cancellation), or a total not updated."),
    ("Change orders add up",
     "The change orders of a contract sum to the change-order total held on the contract.",
     "TOTAL",
     "sum(change_order.change_amount) by change_order.contract_id = shipbuilding_contract.change_order_total by shipbuilding_contract.contract_id",
     "T-TOTAL",
     "The contract total is only trustworthy if it can be rebuilt from the individual amendments.",
     "Requested or rejected change orders are in the list but not in the total, or an amendment was keyed directly into the contract."),
    ("Instalments make up the price",
     "The instalments of a contract add up to its current price.",
     "TOTAL",
     "sum(payment_milestone.milestone_amount) by payment_milestone.contract_id = shipbuilding_contract.contract_price_current by shipbuilding_contract.contract_id",
     "T-TOTAL",
     "The payment plan is a split of the price; every unit of the price falls due at some milestone, typically the largest at delivery.",
     "The plan still reflects the original price and change orders are billed separately, or an instalment is missing."),
    ("Invoices add up to invoiced total",
     "The net amounts of the invoices issued under a contract sum to the invoiced total on the contract.",
     "TOTAL",
     "sum(sales_invoice.net_amount) by sales_invoice.contract_id = shipbuilding_contract.invoiced_total by shipbuilding_contract.contract_id",
     "T-TOTAL",
     "Sales per contract must be the same whether read from the invoice list or from the contract record.",
     "Credit notes or cancelled invoices treated differently, invoices in a currency other than the contract currency, or an invoice booked to the wrong contract."),
    ("Invoice gross amount",
     "Gross amount is net amount plus tax.",
     "IDENTITY",
     "sales_invoice: gross_amount = net_amount + tax_amount",
     "T-MONEY",
     "It is how an invoice is constructed; it also tells us which column is the tax-exclusive figure that counts as sales.",
     "Net and gross columns swapped, tax held elsewhere, or a retention or advance deducted on the invoice."),
    ("Invoice in functional currency",
     "The functional-currency net amount is the invoice net amount times the exchange rate.",
     "IDENTITY",
     "sales_invoice: net_amount_functional = net_amount * fx_rate",
     "T-FX",
     "Vessels are often priced in a currency other than the yard's own; sales can only be added up across invoices in one currency.",
     "Rate stored the other way round, a hedged or contract rate used, or the functional amount is gross while the other is net."),
    ("One contract per vessel",
     "No vessel appears on more than one contract.",
     "UNIQUE",
     "shipbuilding_contract: project_id",
     "T-EXACT",
     "Revenue per ship needs an unambiguous link from the vessel to its price.",
     "A resale after owner default, a contract re-issued under a new number, or a series contract split in an unexpected way."),
    ("Invoice numbers are unique per entity",
     "Within a legal entity, an invoice number is used once.",
     "UNIQUE",
     "sales_invoice: entity_id, invoice_number",
     "T-EXACT",
     "Tax law requires a unique consecutive numbering; a duplicate would count sales twice.",
     "The export contains invoice lines rather than invoices, the same file was loaded twice, or number ranges restart each year."),
    ("Contract refers to a vessel",
     "Every vessel named on a contract exists in the vessel project list.",
     "REFERENCE",
     "shipbuilding_contract.project_id -> vessel_project.project_id",
     "T-EXACT",
     "A price without a vessel cannot be compared with any cost.",
     "Contract signed before the project was opened, or different numbering between sales and production."),
    ("Invoice refers to a contract",
     "Every contract named on an invoice exists in the contract list.",
     "REFERENCE",
     "sales_invoice.contract_id -> shipbuilding_contract.contract_id",
     "T-EXACT",
     "Newbuilding sales must be traceable to the contract they were billed under.",
     "Repair, conversion or after-sales invoices carrying an order number of another kind, or an incomplete contract list."),
    # --- status and dates ---------------------------------------------------
    ("Delivered vessels have a delivery date",
     "A vessel whose status is delivered has an actual delivery date.",
     "CONDITION",
     "vessel_project: when project_status = STATUS_DELIVERED then delivery_date present",
     "T-EXACT",
     "Handover is a documented legal event (protocol of delivery and acceptance); its date decides in which year a vessel counts as sold or completed.",
     "Status set without the date being maintained, or the date kept only in the contract file."),
    ("Vessels in build have no delivery date",
     "A vessel that is still under construction has no actual delivery date.",
     "CONDITION",
     "vessel_project: when project_status = STATUS_IN_BUILD then delivery_date absent",
     "T-EXACT",
     "It confirms that the column holds the actual and not the planned date, and that the status is kept up to date.",
     "The planned date is stored in the actual-date column, or the status was not updated after handover."),
    ("Approved change orders have an approval date",
     "A change order with status approved carries the date of approval.",
     "CONDITION",
     "change_order: when change_order_status = STATUS_CHANGE_APPROVED then approval_date present",
     "T-EXACT",
     "Only an agreed amendment changes the price; the date fixes the period from which it counts.",
     "Verbal agreement entered as approved, or approval recorded only in correspondence."),
    ("Hull number is unique per yard",
     "Within a legal entity, a hull number identifies one vessel project.",
     "UNIQUE",
     "vessel_project: entity_id, hull_number",
     "T-EXACT",
     "The hull number is the yard's name for a vessel on every drawing, order and timesheet.",
     "Sub-projects (sections, options, warranty) listed as separate projects under one hull, or a cancelled and restarted project."),
    # --- cost ---------------------------------------------------------------
    ("Cost entry in functional currency",
     "The functional-currency amount of a cost posting is the transaction amount times the exchange rate.",
     "IDENTITY",
     "cost_ledger_entry: amount_functional = amount_transaction * fx_rate",
     "T-FX",
     "Major equipment is often bought in foreign currency; the cost of a vessel can only be added up in one currency.",
     "Rate stored the other way round, a hedge rate used, or a revaluation difference posted into the same line."),
    ("Cost entry refers to a vessel",
     "Every vessel named on a cost posting exists in the vessel project list.",
     "REFERENCE",
     "cost_ledger_entry.project_id -> vessel_project.project_id",
     "T-EXACT",
     "A cost that cannot be assigned to a known vessel cannot be part of any vessel's cost.",
     "Cost collected on internal, overhead, repair or warranty orders that share the same ledger."),
    ("Cost ledger adds up to project cost",
     "The cost postings of a vessel sum to the total cost shown in its project summary.",
     "TOTAL",
     "sum(cost_ledger_entry.amount_functional) by cost_ledger_entry.project_id = project_cost_summary.total_cost by project_cost_summary.project_id",
     "T-TOTAL",
     "The reported cost of a vessel must be reproducible from the postings; this is the core of the question what a vessel costs.",
     "Summary at a different cut-off date, postings of both yards in different currencies added together, or overhead added in the summary only."),
    ("Cost categories add up",
     "Total cost of a vessel is the sum of labour, material, subcontract, overhead and other cost.",
     "IDENTITY",
     "project_cost_summary: total_cost = labour_cost_total + material_cost_total + subcontract_cost_total + overhead_cost_total + other_cost_total",
     "T-TOTAL",
     "The breakdown is complete by construction; if it holds, each of the six columns is confirmed to be what it says.",
     "A further category (design, financing, warranty provision, commission) held outside these five."),
    ("Cost at completion",
     "Estimated cost at completion is cost to date plus estimated cost to complete.",
     "IDENTITY",
     "project_cost_summary: estimated_cost_at_completion = total_cost + estimated_cost_to_complete",
     "T-ESTIMATE",
     "This is the definition of the forecast; for a vessel in build, the cost of the vessel is this figure and not the cost to date.",
     "Forecast and actuals taken at different dates, or commitments counted in one figure and not the other."),
    ("Stage of completion",
     "Stage of completion is cost to date divided by estimated cost at completion.",
     "IDENTITY",
     "project_cost_summary: percent_complete = total_cost / estimated_cost_at_completion",
     "T-ESTIMATE",
     "Cost-to-cost is the usual measure of progress for vessels built to order under long-term contract accounting.",
     "Progress measured physically or by milestones, percentage stored on a scale of 100, or uninstalled major equipment excluded from the measure."),
    ("Revenue recognised over time",
     "Revenue recognised to date is the contract revenue times the stage of completion.",
     "IDENTITY",
     "project_cost_summary: revenue_recognised_total = contract_revenue_functional * percent_complete",
     "T-ESTIMATE",
     "Where revenue is recognised over the build period, this is how the figure arises; whether it holds reveals which revenue method the data follows.",
     "Revenue recognised at delivery only, a loss-making contract with revenue capped, or contract revenue converted at a different rate."),
    # --- labour -------------------------------------------------------------
    ("Hours add up",
     "Total hours of a time booking are regular hours plus overtime hours.",
     "IDENTITY",
     "timesheet_entry: total_hours = regular_hours + overtime_hours",
     "T-HOURS",
     "Hours are the main measure of yard output; the total must not hide or double-count overtime.",
     "Further kinds of time (travel, standby, night shift) held in other columns."),
    ("Labour cost of a booking",
     "Labour cost is total hours times the hourly costing rate.",
     "IDENTITY",
     "timesheet_entry: labour_cost = total_hours * hourly_rate",
     "T-MONEY",
     "Own labour reaches the vessel through hours valued at a rate; the identity confirms hours, rate and value together.",
     "Overtime valued with a premium, a rate changed after the fact, or the rate being a sales rate rather than a cost rate."),
    ("Timesheet refers to an employee",
     "Every employee named on a time booking exists in the employee list.",
     "REFERENCE",
     "timesheet_entry.employee_id -> employee.employee_id",
     "T-EXACT",
     "Own labour must come from the yard's own people; this separates it from subcontract labour.",
     "Agency or subcontract workers booking time in the same system, or leavers removed from the master list."),
    ("Timesheet hours add up to project hours",
     "The hours booked to a vessel sum to the labour hours shown in its project summary.",
     "TOTAL",
     "sum(timesheet_entry.total_hours) by timesheet_entry.project_id = project_cost_summary.labour_hours_total by project_cost_summary.project_id",
     "T-TOTAL",
     "Hours per vessel are a key figure in their own right and the basis of the labour cost.",
     "Hours of the other yard or of subcontractors included in the summary, or a different cut-off date."),
    # --- material and purchasing -------------------------------------------
    ("Purchase order line value",
     "Line amount is ordered quantity times unit price.",
     "IDENTITY",
     "purchase_order_line: line_amount = ordered_quantity * unit_price",
     "T-MONEY",
     "Bought-in equipment and material are the largest share of a vessel's cost; the identity confirms quantity, price and value together.",
     "A price per 100 or per tonne with a different quantity unit, a discount or freight on the line, or lump-sum service lines."),
    ("Material issue value",
     "Value of a store issue is issued quantity times unit cost.",
     "IDENTITY",
     "material_issue: issue_value = issued_quantity * unit_cost",
     "T-MONEY",
     "Stock material reaches the vessel at the stock valuation price; the identity confirms quantity, price and value together.",
     "Unit cost shown is today's price rather than the price at issue, or a price unit other than one."),
    # --- general ledger -----------------------------------------------------
    ("Posting documents balance",
     "On every posting document, total debit equals total credit.",
     "IDENTITY",
     "journal_voucher: total_debit = total_credit",
     "T-EXACT",
     "Double-entry bookkeeping: a ledger whose documents do not balance is not a complete ledger.",
     "The export is filtered (only some accounts or one entity), or totals are in different currencies."),
    ("Ledger lines add up to the document",
     "The debit lines of a posting document sum to its total debit.",
     "TOTAL",
     "sum(general_ledger_entry.debit_amount) by general_ledger_entry.journal_id = journal_voucher.total_debit by journal_voucher.journal_id",
     "T-EXACT",
     "It confirms that the line file is complete and that the debit column is read with the right sign.",
     "Lines missing from the export, or negative debits used for reversals and counted differently in the header."),
    ("Ledger line is unique",
     "Posting document and line number together identify one ledger line.",
     "UNIQUE",
     "general_ledger_entry: journal_id, line_number",
     "T-EXACT",
     "Duplicate lines would double sales and cost alike.",
     "Overlapping exports loaded together, or document numbers that restart each year or per entity."),
    ("Ledger line uses a known account",
     "Every account posted to exists in the chart of accounts.",
     "REFERENCE",
     "general_ledger_entry.account_id -> account.account_id",
     "T-EXACT",
     "Whether a line is revenue or cost is known only through its account.",
     "The two entities use different charts of accounts and only one was supplied."),
    # --- group --------------------------------------------------------------
    ("Intercompany charge in both currencies",
     "The buyer's amount is the seller's amount times the exchange rate.",
     "IDENTITY",
     "intercompany_transaction: amount_buyer = amount_seller * fx_rate",
     "T-FX",
     "Work of one yard for a vessel of the other must appear at the same value on both sides, or group cost and group sales cannot be reconciled.",
     "The two entities booked at different rates or dates, or the buyer's amount includes duties or freight."),
]

SEC7_INTRO = (
    "The rules above can show whether the data is consistent. They cannot say which "
    "of several consistent readings is the one the company means. The following "
    "questions need an answer from a person before the three business questions "
    "can be answered at all:"
)
SEC7 = [
    "Which cost categories count as the cost of building a vessel: direct cost only, or also allocated yard overhead, design and engineering, financing cost, guarantees and commissions, warranty work after delivery?",
    "Is the cost of a vessel still in build the cost to date or the estimated cost at completion?",
    "What does \"sales\" mean: revenue recognised under the accounting standard in use (over the build period or at delivery), amounts invoiced, cash received, or the contract value of vessels delivered in the year?",
    "What is \"revenue per ship\": the current contract price including approved change orders, the revenue recognised so far, or the amount invoiced? Do unapproved claims, penalties for late delivery and price adjustment clauses count?",
    "What does \"last year\" mean: calendar or financial year, and by which date is a record assigned to it (posting date, invoice date, delivery date)?",
    "In which currency are group figures stated, and at which rate are the subsidiary's figures converted (rate of each transaction, average rate of the year, closing rate)?",
    "How is work between the two yards treated: at the charged transfer price including any margin, or at the delivering yard's own cost; and are intercompany sales removed from group sales?",
    "How are costs shared between sister vessels of a series (design, jigs, first-of-class effects) assigned to the individual vessel?",
    "Are subsidies, grants and insurance recoveries deducted from cost, added to revenue, or shown separately?",
    "How are cancelled contracts, vessels built for stock or resale, and non-newbuilding business (repair, conversion, service) treated in each of the three questions?",
    "Which system is authoritative when two sources that should agree do not (project summary against cost ledger, contract record against invoices)?",
]

# --------------------------------------------------------------------------
# Document building helpers
# --------------------------------------------------------------------------
def _shade(cell, fill="D9E2F3"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)
    cant = OxmlElement("w:cantSplit")
    cant.set(qn("w:val"), "true")
    trPr.append(cant)


def _set_cell(cell, text, size, bold=False, mono=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if mono:
        run.font.name = "Consolas"
        rPr = run._r.get_or_add_rPr()
        rf = rPr.find(qn("w:rFonts"))
        rf.set(qn("w:cs"), "Consolas")
        rf.set(qn("w:hAnsi"), "Consolas")


def add_table(doc, header, rows, widths_cm, size=9, mono_cols=()):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.autofit = False
    for i, h in enumerate(header):
        c = table.rows[0].cells[i]
        _set_cell(c, h, size, bold=True)
        _shade(c)
    _repeat_header(table.rows[0])
    for r in rows:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            _set_cell(cells[i], v, size, mono=(i in mono_cols))
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    return table


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.5))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.core_properties.title = "Foundation document: shipbuilder building vessels to order"
    doc.core_properties.author = "Controller (general industry knowledge)"

    doc.add_paragraph("Foundation document: shipbuilder building vessels to order", style="Title")
    doc.add_paragraph(
        "What must be true of the data if it is what it claims to be. Written from "
        "general shipbuilding and accounting practice, before any company file was seen."
    )

    # 1
    doc.add_heading(HEADINGS[0], level=1)
    for p in PURPOSE:
        doc.add_paragraph(p)

    # 2
    doc.add_heading(HEADINGS[1], level=1)
    add_table(doc, H_OBJECTS, OBJECTS, [5.0, 13.7, 8.0], mono_cols=(0,))

    # 3
    doc.add_heading(HEADINGS[2], level=1)
    add_table(doc, H_QUANT, QUANTITIES, [5.0, 5.7, 2.5, 13.5], mono_cols=(0, 1, 2))

    # 4
    doc.add_heading(HEADINGS[3], level=1)
    for p in TOL_PROSE:
        doc.add_paragraph(p)
    add_table(doc, H_TOL, TOLERANCES, [2.6, 8.6, 2.0, 2.0, 11.5], mono_cols=(0, 2, 3))
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Minimum share").bold = True
    doc.add_paragraph(MIN_SHARE_PROSE)
    p = doc.add_paragraph()
    r = p.add_run(f"MIN_SHARE = {MIN_SHARE}")
    r.font.name = "Consolas"

    # 5
    doc.add_heading(HEADINGS[4], level=1)
    doc.add_paragraph(
        "Only the company can supply these values. The Example column shows the kind "
        "of value expected; replace it with your own."
    )
    add_table(doc, H_PARAM, PARAMETERS, [6.5, 15.2, 5.0], mono_cols=(0,))

    # 6
    doc.add_heading(HEADINGS[5], level=1)
    doc.add_paragraph(
        "Kind, Formal and Tolerance are read by the checking tool. In the last "
        "column write Yes, No or Partly."
    )
    rows = [(f"R{i:02d}", n, s, k, f, t, w, e, "")
            for i, (n, s, k, f, t, w, e) in enumerate(RULES, start=1)]
    add_table(doc, H_RULES, rows, [1.0, 2.6, 3.9, 1.9, 5.6, 1.8, 4.0, 4.0, 1.9],
              size=8, mono_cols=(4,))

    # 7
    doc.add_heading(HEADINGS[6], level=1)
    doc.add_paragraph(SEC7_INTRO)
    for item in SEC7:
        doc.add_paragraph(item, style="List Bullet")

    doc.save(OUT)


# --------------------------------------------------------------------------
# Verification: re-open the written file and check it against the contract
# --------------------------------------------------------------------------
IDENT = r"[a-z][a-z0-9_]*"
PARAM = r"[A-Z][A-Z0-9_]*"


def _rows(table):
    return [[c.text.strip() for c in row.cells] for row in table.rows]


def verify():
    errors = []
    doc = Document(OUT)

    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    if h1 != HEADINGS:
        errors.append(f"headings differ: {h1}")

    tables = {tuple(_rows(t)[0]): _rows(t)[1:] for t in doc.tables}
    for hdr in (H_OBJECTS, H_QUANT, H_TOL, H_PARAM, H_RULES):
        if tuple(hdr) not in tables:
            errors.append(f"missing table with header {hdr}")
    if errors:
        return errors, {}
    if len(doc.tables) != 5:
        errors.append(f"expected 5 tables, found {len(doc.tables)}")

    snake = re.compile(rf"^{IDENT}$")
    objects = [r[0] for r in tables[tuple(H_OBJECTS)]]
    if len(set(objects)) != len(objects):
        errors.append("duplicate object names")
    for o in objects:
        if not snake.match(o):
            errors.append(f"object not snake_case: {o}")

    quants = {}
    for o, q, t, _m in tables[tuple(H_QUANT)]:
        if o not in objects:
            errors.append(f"quantity on unknown object: {o}.{q}")
        if not snake.match(q):
            errors.append(f"quantity not snake_case: {q}")
        if t not in TYPES:
            errors.append(f"bad type {t} for {o}.{q}")
        if q in quants.setdefault(o, {}):
            errors.append(f"duplicate quantity {o}.{q}")
        quants[o][q] = t
    for o in objects:
        if o not in quants:
            errors.append(f"object without quantities: {o}")

    classes = set()
    for c, _a, ab, rel, _r in tables[tuple(H_TOL)]:
        if not re.match(r"^T-[A-Z]+$", c):
            errors.append(f"bad tolerance class id {c}")
        for v in (ab, rel):
            if v != "-":
                try:
                    float(v)
                except ValueError:
                    errors.append(f"tolerance value not numeric: {c} {v}")
        classes.add(c)
    if "T-EXACT" not in classes:
        errors.append("T-EXACT missing")

    ms = [p.text for p in doc.paragraphs if re.match(r"^MIN_SHARE = 0\.\d+$", p.text.strip())]
    if len(ms) != 1 or not (0 < float(ms[0].split("=")[1]) < 1):
        errors.append(f"MIN_SHARE line problem: {ms}")

    params = [r[0] for r in tables[tuple(H_PARAM)]]
    for p in params:
        if not re.match(rf"^{PARAM}$", p):
            errors.append(f"parameter not UPPER_SNAKE: {p}")

    def has(o, q, where):
        if o not in quants:
            errors.append(f"{where}: unknown object {o}")
            return False
        if q not in quants[o]:
            errors.append(f"{where}: unknown quantity {o}.{q}")
            return False
        return True

    used = set()
    rules = tables[tuple(H_RULES)]
    for i, (rid, name, stmt, kind, formal, tol, why, exc, applies) in enumerate(rules, start=1):
        if rid != f"R{i:02d}":
            errors.append(f"bad id sequence at {rid}")
        if not all((name, stmt, why, exc)):
            errors.append(f"{rid}: empty text column")
        if applies:
            errors.append(f"{rid}: 'Applies to us' not empty")
        if tol not in classes:
            errors.append(f"{rid}: unknown tolerance {tol}")
        if "\n" in formal:
            errors.append(f"{rid}: formal not on one line")
        if kind not in KINDS:
            errors.append(f"{rid}: unknown kind {kind}")
            continue

        if kind == "IDENTITY":
            m = re.match(rf"^({IDENT}): ([^=]+) = ([^=]+)$", formal)
            if not m:
                errors.append(f"{rid}: IDENTITY syntax")
                continue
            o = m.group(1)
            for side in (m.group(2), m.group(3)):
                if not re.match(r"^[a-z0-9_+\-*/(). ]+$", side):
                    errors.append(f"{rid}: illegal characters in expression")
                for q in re.findall(IDENT, side):
                    if has(o, q, rid):
                        used.add((o, q))
                try:  # expression must be well-formed arithmetic
                    compile(re.sub(IDENT, "1", side).strip(), rid, "eval")
                except SyntaxError:
                    errors.append(f"{rid}: malformed expression '{side}'")
        elif kind == "UNIQUE":
            m = re.match(rf"^({IDENT}): ({IDENT}(?:, {IDENT})*)$", formal)
            if not m:
                errors.append(f"{rid}: UNIQUE syntax")
                continue
            for q in m.group(2).split(", "):
                if has(m.group(1), q, rid):
                    used.add((m.group(1), q))
        elif kind == "REFERENCE":
            m = re.match(rf"^({IDENT})\.({IDENT}) -> ({IDENT})\.({IDENT})$", formal)
            if not m:
                errors.append(f"{rid}: REFERENCE syntax")
                continue
            for o, q in ((m.group(1), m.group(2)), (m.group(3), m.group(4))):
                if has(o, q, rid):
                    used.add((o, q))
        elif kind == "CONDITION":
            m = re.match(rf"^({IDENT}): when ({IDENT}) = ({PARAM}) then ({IDENT}) (present|absent)$", formal)
            if not m:
                errors.append(f"{rid}: CONDITION syntax")
                continue
            o, sq, par, q = m.group(1), m.group(2), m.group(3), m.group(4)
            if has(o, sq, rid):
                used.add((o, sq))
                if quants[o][sq] != "status":
                    errors.append(f"{rid}: {o}.{sq} is not of type status")
            if has(o, q, rid):
                used.add((o, q))
            if par not in params:
                errors.append(f"{rid}: unknown parameter {par}")
        elif kind == "TOTAL":
            m = re.match(
                rf"^sum\(({IDENT})\.({IDENT})\) by ({IDENT})\.({IDENT}) = "
                rf"({IDENT})\.({IDENT}) by ({IDENT})\.({IDENT})$", formal)
            if not m:
                errors.append(f"{rid}: TOTAL syntax")
                continue
            g = m.groups()
            if g[0] != g[2] or g[4] != g[6]:
                errors.append(f"{rid}: key must belong to the same object as the figure")
            for o, q in ((g[0], g[1]), (g[2], g[3]), (g[4], g[5]), (g[6], g[7])):
                if has(o, q, rid):
                    used.add((o, q))

    stats = {
        "objects": len(objects),
        "quantities": sum(len(v) for v in quants.values()),
        "tolerance_classes": len(classes),
        "parameters": len(params),
        "rules": len(rules),
        "min_share": ms[0] if ms else None,
        "kinds": {k: sum(1 for r in rules if r[3] == k) for k in sorted(KINDS)},
        "objects_without_rule": sorted(set(objects) - {o for o, _ in used}),
        "tolerances_unused": sorted(classes - {r[5] for r in rules}),
    }
    return errors, stats


if __name__ == "__main__":
    build()
    errs, stats = verify()
    print(f"written: {OUT}")
    for k, v in stats.items():
        print(f"  {k}: {v}")
    if errs:
        print("VERIFICATION FAILED")
        for e in errs:
            print("  -", e)
        sys.exit(1)
    print("verification OK")
