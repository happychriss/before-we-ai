#!/usr/bin/env python3
"""Generator for the affiliate landscape (Meridia Pharma Norvania d.o.o., FY2025).

DO NOT READ if you are implementing or testing against this landscape: this
file places the seeded cases.

How it works
------------
1. `build_true_world()` builds what really happened: master data, purchase
   orders, receipts, invoices, payments and the *correct* manual journals.
   Real-but-awkward business events (kind 1) are part of this world.
2. `simulate()` plays a world through a small model of the local ERP and
   returns the tables (subledgers and ledger). The simulation of the true world
   is the CLEAN BASE; `selfcheck.run_checks` must pass on it.
3. `apply_disorder()` turns a deep copy of the true world into the recorded
   world: real errors (kind 2) and undecidable things (kind 3).
4. The recorded world is simulated and written to `data/` with ordinary noise.
   The answer key is computed from both worlds and written to `answer-key/`.

Deterministic: fixed seed, fixed timestamps in the workbook and the PDFs.

    .venv/bin/python corpora/affiliate/generator/generate_affiliate.py
"""

from __future__ import annotations

import copy
import csv
import datetime as dt
import hashlib
import io
import random
import re
import sys
import zipfile
from collections import defaultdict
from fractions import Fraction as F
from pathlib import Path
from types import SimpleNamespace as NS

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))

SEED = 20251003
D = dt.date
CUT = D(2025, 7, 1)
YS = D(2025, 1, 1)
YE = D(2025, 12, 31)
DAY = dt.timedelta(days=1)

EOM = {(2024, 10): "11.1820", (2024, 11): "11.2135", (2024, 12): "11.2410",
       (2025, 1): "11.2890", (2025, 2): "11.3325", (2025, 3): "11.3105",
       (2025, 4): "11.3840", (2025, 5): "11.4420", (2025, 6): "11.4915",
       (2025, 7): "11.5330", (2025, 8): "11.6045", (2025, 9): "11.5810",
       (2025, 10): "11.6720", (2025, 11): "11.7435", (2025, 12): "11.8260"}
CLOSING = F(EOM[(2025, 12)])

GRIR = ("2810", "2815")
EXTRACT_ACCTS = {"1400", "1410", "1450", "2810", "2815", "3300", "3310", "3350", "5890", "6950"}


def eom_rate(y, m):
    return F(EOM[(y, m)])


def book(d):
    y, m = (d.year, d.month - 1) if d.month > 1 else (d.year - 1, 12)
    return F(EOM[(y, m)])


def rnd(x):
    return int(round(F(x)))


def bd(d):
    while d.weekday() >= 5:
        d += DAY
    return d


def month_end_bd(d):
    n = D(d.year + (d.month == 12), d.month % 12 + 1, 1) - DAY
    while n.weekday() >= 5:
        n -= DAY
    return n


def next_thursday(d):
    while d.weekday() != 3:
        d += DAY
    return d


def money(c):
    s = "-" if c < 0 else ""
    c = abs(int(c))
    return f"{s}{c // 100}.{c % 100:02d}"


def qfmt(q):
    q = F(q)
    if q.denominator == 1:
        return str(q.numerator)
    return f"{float(q):.3f}".rstrip("0").rstrip(".")


# --------------------------------------------------------------- master data

VENDORS = [
    # no, name, city, ctry, cur, terms, days, cat, reffmt, bills_in_base_unit
    ("79001", "Meridia Supply Chain B.V.", "Rotterdam", "NL", "EUR", "IC60", 60, "IC", 9, False),
    ("70011", "Kartonaza Velmir d.o.o.", "Velograd", "NV", "NVK", "N30", 30, "PACK", 0, False),
    ("70012", "Brenko Pak d.o.o.", "Port Sorna", "NV", "NVK", "N30", 30, "PACK", 1, False),
    ("70013", "Etiketa Sorvin d.o.o.", "Kresnik", "NV", "NVK", "N30", 30, "PACK", 2, True),
    ("70014", "Tisk Lumar d.o.o.", "Velograd", "NV", "NVK", "N14", 14, "PACK", 3, True),
    ("70015", "Folia Trand d.o.o.", "Malenburg", "NV", "NVK", "N30", 30, "PACK", 4, False),
    ("70016", "Ambalaza Krestov d.d.", "Dolna Tirna", "NV", "NVK", "N45", 45, "PACK", 0, False),
    ("70021", "Labtek Soran d.o.o.", "Velograd", "NV", "NVK", "N30", 30, "LAB", 1, False),
    ("70022", "Reagens Mivor d.o.o.", "Haldova", "NV", "NVK", "N30", 30, "LAB", 2, True),
    ("70023", "Analitika Preles d.o.o.", "Velograd", "NV", "NVK", "N30", 30, "LAB", 3, False),
    ("70024", "Vitrum Laborbedarf Halden GmbH", "Graz", "AT", "EUR", "N30", 30, "LAB", 4, False),
    ("70025", "Sciomed Instruments Kereven GmbH", "Jena", "DE", "EUR", "N30", 30, "LAB", 0, False),
    ("70026", "Medilab Torvan d.o.o.", "Port Sorna", "NV", "NVK", "ADV", 0, "LAB", 1, False),
    ("70031", "Papirnica Osrek d.o.o.", "Velograd", "NV", "NVK", "N14", 14, "OFF", 2, True),
    ("70032", "Biro Dalven d.o.o.", "Kresnik", "NV", "NVK", "N30", 30, "OFF", 3, False),
    ("70033", "Toner Servis Plavic d.o.o.", "Velograd", "NV", "NVK", "N30", 30, "OFF", 4, False),
    ("70034", "Uredska Oprema Hrist d.o.o.", "Malenburg", "NV", "NVK", "N30", 30, "OFF", 0, False),
    ("70041", "Paleta Drovnik d.o.o.", "Dolna Tirna", "NV", "NVK", "N30", 30, "LOG", 1, False),
    ("70042", "Coldbox Systems Fennvik B.V.", "Venlo", "NL", "EUR", "N45", 45, "LOG", 2, False),
    ("70043", "Termo Log Savrin d.o.o.", "Port Sorna", "NV", "NVK", "N30", 30, "LOG", 3, False),
    ("70044", "Logistika Materijal Benko d.o.o.", "Haldova", "NV", "NVK", "ADV", 0, "LOG", 4, False),
    ("70045", "Senzor Tehnika Olvar d.o.o.", "Velograd", "NV", "NVK", "N30", 30, "LOG", 0, False),
    # service vendors: invoices without purchase order
    ("70051", "Nepremicnine Castrum d.o.o.", "Velograd", "NV", "NVK", "N14", 14, "SRV", 1, False),
    ("70052", "Energija Norvania d.d.", "Velograd", "NV", "NVK", "N14", 14, "SRV", 2, False),
    ("70053", "Norvatel d.d.", "Velograd", "NV", "NVK", "N14", 14, "SRV", 4, False),
    ("70054", "Autolease Vendra d.o.o.", "Port Sorna", "NV", "NVK", "N14", 14, "SRV", 3, False),
    ("70055", "Cisto Servis Malen d.o.o.", "Malenburg", "NV", "NVK", "N30", 30, "SRV", 0, False),
    ("70056", "Varnost Tregal d.o.o.", "Velograd", "NV", "NVK", "N30", 30, "SRV", 1, False),
    ("70057", "Konsalt Ribar d.o.o.", "Velograd", "NV", "NVK", "N30", 30, "SRV", 2, False),
    ("70058", "Revizija Holmek d.o.o.", "Velograd", "NV", "NVK", "N30", 30, "SRV", 3, False),
    ("70059", "Odvetniki Strelan in partnerji o.p.", "Velograd", "NV", "NVK", "N30", 30, "SRV", 4, False),
    ("70060", "Agencija Pikto d.o.o.", "Kresnik", "NV", "NVK", "N30", 30, "SRV", 0, False),
    ("70061", "Hitra Dostava Norvania d.o.o.", "Velograd", "NV", "NVK", "N14", 14, "SRV", 1, False),
    ("70062", "Kadri Plus Dorvan d.o.o.", "Velograd", "NV", "NVK", "N14", 14, "SRV", 2, False),
    ("70063", "Kongres Servis Lipan d.o.o.", "Dolna Tirna", "NV", "NVK", "N30", 30, "SRV", 3, False),
    ("70064", "Prevodi Semra s.p.", "Haldova", "NV", "NVK", "N14", 14, "SRV", 4, False),
    ("70065", "Spedicija Karvel d.o.o.", "Port Sorna", "NV", "NVK", "N30", 30, "SRV", 0, False),
    ("70066", "Sistemi Datren d.o.o.", "Velograd", "NV", "NVK", "N30", 30, "SRV", 1, False),
    ("70067", "Farma Odpad Rekol d.o.o.", "Malenburg", "NV", "NVK", "N30", 30, "SRV", 2, False),
    ("70068", "Zavarovalnica Tirna d.d.", "Dolna Tirna", "NV", "NVK", "N30", 30, "SRV", 3, False),
]

# vendor no -> [(descr, base unit, order unit, conv, price per order unit, qmin, qmax)]
TP_MATERIALS = {
    "70011": [("Shipping carton 400x300x250", "PCE", "BDL", 25, 455.00, 40, 200),
              ("Shipping carton 600x400x300", "PCE", "BDL", 25, 690.00, 30, 150),
              ("Carton divider set", "PCE", "BDL", 50, 388.00, 50, 200),
              ("Pallet box heavy duty", "PCE", "PCE", 1, 1460.00, 20, 100)],
    "70012": [("Bubble wrap 1.2 m", "ROL", "CAR", 6, 412.00, 50, 300),
              ("Paper filler 70 g", "KG", "BAG", 15, 655.00, 20, 120),
              ("Packing tape 48 mm", "ROL", "CAR", 36, 472.00, 20, 100)],
    "70013": [("Label thermal 100x60", "PCE", "ROL", 1000, 61.50, 200, 1200),
              ("Label cold-chain warning", "PCE", "ROL", 500, 96.00, 50, 300),
              ("Label serialisation blank 50x30", "PCE", "ROL", 2000, 111.80, 100, 600)],
    "70014": [("Patient leaflet reprint type A", "PCE", "BOX", 500, 890.00, 10, 60),
              ("Patient leaflet reprint type B", "PCE", "BOX", 500, 935.00, 10, 60),
              ("Outer sticker NVK price", "PCE", "BOX", 2000, 640.00, 20, 80)],
    "70015": [("Stretch film 500 mm 23 my", "ROL", "CAR", 6, 318.00, 30, 200),
              ("Shrink hood pallet", "PCE", "BDL", 50, 742.50, 10, 80)],
    "70016": [("Insulated shipper 12 l", "PCE", "PCE", 1, 268.00, 20, 150),
              ("Gel pack 500 g", "PCE", "CAR", 24, 205.00, 40, 400),
              ("Thermal blanket pallet", "PCE", "PCE", 1, 1180.00, 5, 40)],
    "70021": [("Pipette tips 200 ul", "PCE", "CS", 960, 980.00, 10, 60),
              ("Nitrile gloves M", "BOX", "CAR", 10, 620.00, 10, 50),
              ("Sample vials 2 ml", "PCE", "BOX", 100, 2350.00, 5, 30)],
    "70022": [("Buffer solution pH 7", "BTL", "BOX", 6, 1275.00, 5, 40),
              ("Reference standard kit", "PCE", "PCE", 1, 3480.00, 2, 20),
              ("Solvent HPLC grade 2.5 l", "BTL", "BOX", 4, 1384.20, 5, 30)],
    "70023": [("HPLC column C18", "PCE", "PCE", 1, 14900.00, 1, 6),
              ("Filter membranes 0.45", "PCE", "BOX", 100, 1126.40, 5, 30),
              ("Syringe filters", "PCE", "BOX", 50, 890.00, 5, 40)],
    "70024": [("Volumetric flask set", "SET", "SET", 1, 660.00, 2, 12),
              ("Glass burette 50 ml", "PCE", "PCE", 1, 184.00, 2, 20),
              ("Desiccator 300 mm", "PCE", "PCE", 1, 412.00, 1, 6)],
    "70025": [("Analytical balance module", "PCE", "PCE", 1, 7450.00, 1, 3),
              ("Calibration weight set", "SET", "SET", 1, 660.00, 1, 8),
              ("Stability chamber sensor", "PCE", "PCE", 1, 1290.00, 1, 6)],
    "70026": [("Microbiology test kit", "KIT", "KIT", 1, 2140.00, 3, 20),
              ("Culture media plates", "PCE", "BOX", 20, 468.00, 10, 60)],
    "70031": [("Copy paper A4 80 g", "REAM", "BOX", 5, 310.00, 5, 40),
              ("Envelopes C4", "PCE", "BOX", 250, 268.00, 5, 30)],
    "70032": [("Archive box", "PCE", "BDL", 20, 205.00, 5, 40),
              ("Ring binder A4", "PCE", "CAR", 20, 535.00, 3, 20)],
    "70033": [("Toner cartridge black", "PCE", "PCE", 1, 1090.00, 2, 15),
              ("Toner cartridge colour set", "SET", "SET", 1, 2760.00, 1, 8)],
    "70034": [("Desk organiser set", "SET", "SET", 1, 315.00, 5, 30),
              ("Whiteboard markers", "PCE", "BOX", 10, 142.00, 5, 40)],
    "70041": [("EUR pallet new", "PCE", "PCE", 1, 388.00, 50, 300),
              ("Pallet collar", "PCE", "PCE", 1, 460.00, 20, 120)],
    "70042": [("Cold box 60 l validated", "PCE", "PCE", 1, 318.00, 10, 60),
              ("PCM panel set", "SET", "SET", 1, 96.50, 20, 150)],
    "70043": [("Cool container liner", "PCE", "PCE", 1, 1090.00, 50, 300),
              ("Transport seal numbered", "PCE", "BOX", 500, 1220.00, 10, 80)],
    "70044": [("Roll container", "PCE", "PCE", 1, 3350.00, 2, 12),
              ("Load securing strap", "PCE", "CAR", 10, 640.00, 5, 30)],
    "70045": [("Temperature data logger", "PCE", "PCE", 1, 1490.00, 10, 100),
              ("Logger docking station", "PCE", "PCE", 1, 2140.00, 1, 10),
              ("Humidity indicator card", "PCE", "BOX", 200, 890.00, 2, 20)],
}
GRP_PREFIX = {"FERT": 100, "PACK": 200, "LAB": 300, "OFF": 400, "LOG": 500}

BRANDS = ["Meriflux", "Meritan", "Mericard", "Merizol", "Merilip",
          "Merigast", "Merivent", "Merinor", "Meripain", "Meriderm"]
FORMS = [("5 mg tbl 28x", 48), ("10 mg tbl 28x", 48), ("20 mg tbl 56x", 24), ("40 mg caps 30x", 24),
         ("100 mg/5 ml susp 100 ml", 12), ("250 mg inj 5 amp", 12), ("crm 30 g", 24),
         ("inhal 120 dos", 12)]

SERVICES = [
    # vendor, gl, cc, text, monthly?, amount range (NVK net)
    ("70051", "6100", "9100", "Office rent", True, (86000, 86000)),
    ("70052", "6110", "9100", "Electricity and heating", True, (9000, 21000)),
    ("70053", "6410", "9100", "Telephony and data", True, (6200, 7400)),
    ("70054", "6200", "9300", "Car fleet leasing", True, (41500, 41500)),
    ("70055", "6990", "9100", "Office cleaning", True, (7800, 7800)),
    ("70056", "6990", "9400", "Warehouse security", True, (12400, 12400)),
    ("70057", "6600", "9200", "Consulting pricing dossier", False, (18000, 95000)),
    ("70058", "6600", "9200", "Audit fee instalment", False, (60000, 140000)),
    ("70059", "6610", "9200", "Legal advice", False, (8000, 46000)),
    ("70060", "6500", "9300", "Promotional materials campaign", False, (15000, 120000)),
    ("70061", "6700", "9400", "Courier services", False, (2500, 14000)),
    ("70062", "6710", "9400", "Temporary staff warehouse", False, (14000, 52000)),
    ("70063", "6510", "9300", "Congress participation", False, (20000, 160000)),
    ("70064", "6990", "9300", "Translation leaflet texts", False, (1800, 9000)),
    ("70065", "6700", "9400", "Customs clearance and haulage", False, (6000, 38000)),
    ("70066", "6990", "9100", "IT support hours", False, (5000, 30000)),
    ("70067", "6990", "9400", "Disposal of expired goods", False, (4000, 22000)),
    ("70068", "6990", "9100", "Insurance premium", False, (30000, 88000)),
]

ACCOUNTS = [
    # local, name, group v1, group v2
    ("1000", "Cash on hand", "100100", "100100"),
    ("1100", "Bank NVK current account", "100200", "100200"),
    ("1110", "Bank EUR account", "100200", "100210"),
    ("1200", "Trade receivables pharmacies", "110100", "110100"),
    ("1210", "Trade receivables wholesalers", "110100", "110100"),
    ("1250", "Allowance for doubtful receivables", "110900", "110900"),
    ("1300", "Receivables group companies", "115000", "115000"),
    ("1400", "Inventory trade goods", "120100", "120100"),
    ("1410", "Inventory other materials and supplies", "120300", "120300"),
    ("1450", "Goods in transit", "120500", "120500"),
    ("1490", "Inventory allowance", "120900", "120900"),
    ("1500", "Prepaid expenses", "130100", "130100"),
    ("1550", "Advances to suppliers", "130200", "130200"),
    ("1570", "Input VAT", "135000", "135000"),
    ("1580", "VAT receivable", "135000", "135000"),
    ("1600", "Other receivables", "139000", "139000"),
    ("1700", "Furniture and equipment", "150100", "150100"),
    ("1710", "Vehicles", "150200", "150200"),
    ("1790", "Accumulated depreciation", "150900", "150900"),
    ("1800", "Software", "155100", "155100"),
    ("2000", "Share capital", "300100", "300100"),
    ("2100", "Retained earnings", "310100", "310100"),
    ("2500", "Provision for returns", "250100", "250100"),
    ("2510", "Provision for bonuses", "250200", "250200"),
    ("2600", "Accrued expenses", "240100", "240100"),
    ("2810", "GR/IR clearing", "211500", "211510"),
    ("2815", "GR/IR clearing group companies", None, "211520"),
    ("3300", "Trade payables domestic", "210100", "210100"),
    ("3310", "Trade payables foreign", "210100", "210150"),
    ("3350", "Payables group companies", "215000", "215000"),
    ("3400", "Payroll payables", "220100", "220100"),
    ("3410", "Social security payable", "220200", "220200"),
    ("3500", "Output VAT", "225000", "225000"),
    ("3510", "VAT payable", "225000", "225000"),
    ("3600", "Corporate income tax payable", "226000", "226000"),
    ("3700", "Other payables", "229000", "229000"),
    ("4000", "Sales wholesalers", "400100", "400100"),
    ("4010", "Sales hospitals", "400100", "400150"),
    ("4050", "Sales deductions and rebates", "405000", "405000"),
    ("4100", "Other operating income", "410000", "410000"),
    ("5000", "Cost of goods sold", "500100", "500100"),
    ("5100", "Inventory write-downs", "500500", "500500"),
    ("5200", "Freight in", "500300", "500300"),
    ("5890", "Price differences purchasing", "500900", "500910"),
    ("5895", "Inventory count differences", "500900", "500920"),
    ("6000", "Salaries", "600100", "600100"),
    ("6010", "Social contributions", "600200", "600200"),
    ("6100", "Rent", "610100", "610100"),
    ("6110", "Utilities", "610200", "610200"),
    ("6200", "Vehicle costs", "620100", "620100"),
    ("6300", "Travel", "630100", "630100"),
    ("6400", "Office supplies", "640100", "640100"),
    ("6410", "Telecommunication", "640200", "640200"),
    ("6500", "Marketing and promotion", "650100", "650100"),
    ("6510", "Medical congresses", "650200", "650200"),
    ("6600", "Consulting and audit", "660100", "660100"),
    ("6610", "Legal fees", "660100", "660200"),
    ("6700", "Logistics services", "670100", "670100"),
    ("6710", "Temporary staff", "600300", "600300"),
    ("6800", "Depreciation", "680100", "680100"),
    ("6950", "Exchange differences", "690100", "690110"),
    ("6990", "Other operating expenses", "699000", "699000"),
    ("7000", "Interest", "700100", "700100"),
    ("8000", "Income tax", "800100", "800100"),
]
GROUP_ACCTS = {
    "100100": "Cash", "100200": "Bank balances", "100210": "Bank balances foreign currency",
    "110100": "Trade receivables third parties", "110900": "Allowance trade receivables",
    "115000": "Receivables group companies", "120100": "Inventories trade goods",
    "120150": "Trade goods in transit from group companies",
    "120300": "Inventories materials and supplies", "120500": "Goods in transit",
    "120900": "Inventory allowances", "130100": "Prepaid expenses", "130200": "Advances to suppliers",
    "135000": "VAT receivable", "139000": "Other receivables", "150100": "Equipment",
    "150200": "Vehicles", "150900": "Accumulated depreciation", "155100": "Software",
    "210100": "Trade payables third parties", "210150": "Trade payables third parties foreign",
    "211500": "Goods received not invoiced",
    "211510": "Goods received not invoiced third parties",
    "211520": "Goods received not invoiced group companies",
    "215000": "Payables group companies", "220100": "Payroll liabilities",
    "220200": "Social security liabilities", "225000": "VAT payable", "226000": "Income tax payable",
    "229000": "Other liabilities", "240100": "Accrued expenses", "250100": "Provision returns",
    "250200": "Provision bonuses", "300100": "Share capital", "310100": "Retained earnings",
    "400100": "Net sales third parties", "400150": "Net sales hospitals and institutions",
    "405000": "Sales deductions", "410000": "Other operating income", "500100": "Cost of sales",
    "500300": "Inbound freight", "500500": "Inventory write-downs",
    "500900": "Purchase price and inventory differences", "500910": "Purchase price differences",
    "500920": "Inventory differences", "600100": "Salaries and wages", "600200": "Social charges",
    "600300": "External staff", "610100": "Rent", "610200": "Utilities", "620100": "Vehicle costs",
    "630100": "Travel", "640100": "Office costs", "640200": "Communication",
    "650100": "Promotion", "650200": "Congresses and events", "660100": "Professional fees",
    "660200": "Legal fees", "670100": "Distribution services", "680100": "Depreciation",
    "690100": "Exchange gains and losses", "690110": "Exchange differences operating",
    "699000": "Other operating expenses", "700100": "Interest", "800100": "Income tax",
}


class World:
    def __init__(self):
        self.vendors = {}
        self.mats = {}
        self.pos = []
        self.ev = []
        self.seq = 0
        self.tags = {}
        self.k1 = defaultdict(list)       # kind-1 catalogue: category -> [(po key, pos)]
        self.k7 = []                      # lines whose promised date was overwritten
        self.mapping = []                 # rows of the mapping sheet
        self.refs = set()

    def nseq(self):
        self.seq += 1
        return self.seq


def build_master(W, rng):
    for no, name, city, ctry, cur, terms, days, cat, fmt, base in VENDORS:
        W.vendors[no] = NS(no=no, name=name, city=city, ctry=ctry, cur=cur, terms=terms, days=days,
                           cat=cat, fmt=fmt, base=base, mats=[], refbase=rng.randint(1000, 60000))
    n = defaultdict(int)
    for i, brand in enumerate(BRANDS):
        for j in range(3):
            form, conv = FORMS[(i * 3 + j) % len(FORMS)]
            n["FERT"] += 1
            no = f"{GRP_PREFIX['FERT']}{n['FERT']:03d}"
            pack = rng.randint(380, 4600)                   # EUR cents per pack
            W.mats[no] = NS(no=no, descr=f"{brand} {form}", grp="FERT", bu="PK", ou="CAR", conv=conv,
                            std=pack * conv, cur="EUR", qmin=10, qmax=110, vend="79001")
            W.vendors["79001"].mats.append(no)
    for vno, rows in TP_MATERIALS.items():
        v = W.vendors[vno]
        for descr, bu, ou, conv, price, qmin, qmax in rows:
            n[v.cat] += 1
            no = f"{GRP_PREFIX[v.cat]}{n[v.cat]:03d}"
            W.mats[no] = NS(no=no, descr=descr, grp=v.cat, bu=bu, ou=ou, conv=conv,
                            std=rnd(F(str(price)) * 100), cur=v.cur, qmin=qmin, qmax=qmax, vend=vno)
            v.mats.append(no)
    for loc, name, g1, g2 in ACCOUNTS:
        if g1:
            W.mapping.append(NS(loc=loc, name=name, grp=g1, vfrom=D(2022, 1, 1), vto=D(2025, 6, 30),
                                by="MHAL"))
        W.mapping.append(NS(loc=loc, name=name, grp=g2, vfrom=CUT, vto=D(9999, 12, 31), by="JKOV"))
    W.mapping.append(NS(loc="6900", name="Exchange differences (old)", grp="690100",
                        vfrom=D(2020, 1, 1), vto=D(2023, 12, 31), by="MHAL"))


# ------------------------------------------------------------ event builders

def mat_by(W, vno, idx):
    return W.vendors[vno].mats[idx]


def new_po(W, vno, date, lines, lead=14, tag=None, buyer=None, promised=None):
    """lines: [(material no, qty, price cents or None)]"""
    v = W.vendors[vno]
    key = len(W.pos) + 1
    po = NS(key=key, date=date, vend=vno, cur=v.cur, buyer=buyer or ["E01", "E02", "E03"][(key * 5 + date.day) % 3],
            lines=[], tag=tag, no=None,
            crt="LNOV" if v.cat == "IC" else "PSIM")
    prom = promised or bd(date + lead * DAY)
    for i, (m, q, p) in enumerate(lines):
        price = W.mats[m].std if p is None else p
        po.lines.append(NS(pos=(i + 1) * 10, mat=m, qty=q, price0=price, price=price,
                           promised_true=prom, promised_rec=prom if date >= CUT else None, chgd_on=None))
    W.pos.append(po)
    if tag:
        W.tags.setdefault(tag, NS()).po = po
    return po


def line_of(po, pos):
    for l in po.lines:
        if l.pos == pos:
            return l
    raise KeyError(pos)


def add_gr(W, po, post, doc, items, dn, usr="WH01", txt=""):
    """items: {pos: qty}"""
    e = NS(k="GR", seq=W.nseq(), po=po, post=post, doc=doc, dn=dn, usr=usr, txt=txt, no=None,
           items=[NS(pos=p, qty=q, mat=None, mvt="101", rev=None, lc=None) for p, q in items.items()])
    W.ev.append(e)
    return e


def add_rev(W, gr, post, doc=None, qtys=None, usr=None, txt=""):
    """Reverse `gr` (all items in full, or {index: qty})."""
    items = []
    for i, it in enumerate(gr.items):
        q = it.qty if qtys is None else qtys.get(i)
        if q:
            items.append(NS(pos=it.pos, qty=q, mat=it.mat, mvt="102", rev=it, lc=None))
    e = NS(k="GR", seq=W.nseq(), po=gr.po, post=post, doc=doc or gr.doc, dn=gr.dn, usr=usr or gr.usr,
           txt=txt, no=None, items=items, rev_of=gr)
    W.ev.append(e)
    return e


def vref(W, v, date, rng):
    while True:
        n = v.refbase + (date - D(2024, 11, 1)).days * 7 + rng.randint(0, 6)
        y = date.year
        ref = [f"INV-{y}-{n:05d}", f"R/{n}/{y % 100}", f"{y}/{n:05d}", f"F-{n:06d}", f"{n:07d}"][v.fmt] \
            if v.fmt < 9 else f"MSC{90000000 + n}"
        if (v.no, ref) not in W.refs:
            W.refs.add((v.no, ref))
            return ref


def inv_lines(W, po, items, price=None, short=None):
    """Invoice lines for {pos: qty in order units}, in the unit the vendor bills in."""
    v = W.vendors[po.vend]
    out = []
    for pos, q in items.items():
        l = line_of(po, pos)
        m = W.mats[l.mat]
        p = (price or {}).get(pos, l.price0)
        if v.base and m.conv > 1:
            qb = q * m.conv - (short or {}).get(pos, 0)
            out.append(NS(po=po, pos=pos, mat=l.mat, qty=qb, uom=m.bu, net=rnd(F(qb, m.conv) * p),
                          gl="", cc="", txt=""))
        else:
            out.append(NS(po=po, pos=pos, mat=l.mat, qty=q, uom=m.ou, net=rnd(F(q) * p), gl="", cc="",
                          txt=""))
    return out


def add_inv(W, vno, inv_dt, post, lines, rng, typ="RE", ref=None, usr=None, txt="", pay="auto"):
    v = W.vendors[vno]
    if usr is None:
        usr = rng.choices(["TBER", "MHAL", "JKOV" if post <= D(2025, 7, 31) else "DPET"], [84, 9, 7])[0]
    e = NS(k="INV", seq=W.nseq(), vend=vno, cur=v.cur, inv_dt=inv_dt, post=post, lines=lines, typ=typ,
           ref=ref or vref(W, v, inv_dt, rng), usr=usr, txt=txt, pay=pay, no=None,
           due=inv_dt + v.days * DAY, taxrate=F(20, 100) if v.ctry == "NV" else F(0))
    W.ev.append(e)
    return e


def add_pay(W, inv, date, amt=None, cur=None):
    e = NS(k="PAY", seq=W.nseq(), inv=inv, post=date, amt=amt, cur=cur, no=None)
    W.ev.append(e)
    return e


def add_mj(W, post, items, usr=None, txt="GR/IR clearing", tag=None):
    """items: [NS(po, pos, amt=None, factor=1, acct=None)]"""
    if usr is None:
        usr = "JKOV" if post <= D(2025, 7, 31) else "DPET"
    e = NS(k="MJ", seq=W.nseq(), post=post, items=items, usr=usr, txt=txt, lump=False, no=None)
    W.ev.append(e)
    if tag:
        W.tags.setdefault(tag, NS()).mj = e
    return e


def clr(po, pos, **kw):
    d = dict(po=po, pos=pos, amt=None, factor=1, acct=None)
    d.update(kw)
    return NS(**d)


def inv_total(e):
    net = sum(l.net for l in e.lines)
    tax = rnd(net * e.taxrate)
    return net, tax, net + tax


def schedule_payments(W):
    """Weekly run (Thursday) for third parties, monthly netting on the 25th for the group."""
    paid = {id(e.inv) for e in W.ev if e.k == "PAY"}
    runs = defaultdict(list)
    for e in [x for x in W.ev if x.k == "INV" and x.typ == "RE" and x.pay == "auto" and id(x) not in paid]:
        v = W.vendors[e.vend]
        first = max(e.due, e.post + DAY)
        if v.cat == "IC":
            run = bd(D(first.year, first.month, 25)) if first.day <= 25 else \
                bd(D(first.year + (first.month == 12), first.month % 12 + 1, 25))
        else:
            run = next_thursday(first)
        if run <= YE:
            add_pay(W, e, run)
            runs[e.vend].append(run)
    for e in [x for x in W.ev if x.k == "INV" and x.typ == "KG" and x.pay == "auto" and id(x) not in paid]:
        later = sorted(r for r in runs[e.vend] if r > e.post)
        if later:
            add_pay(W, e, later[0])


# ------------------------------------------------------------ the true world

def rand_day(rng, a, b):
    return bd(a + rng.randint(0, (b - a).days) * DAY)


def split_qty(rng, q, n):
    if n == 1 or q < n * 2:
        return [q] + [0] * (n - 1)
    cuts = sorted(rng.sample(range(1, q), n - 1))
    return [b - a for a, b in zip([0] + cuts, cuts + [q])]


def dn_no(rng, v):
    if v.cat == "IC":
        return f"MSC-DN-{rng.randint(410000, 499999)}"
    return rng.choice(["LS", "DN", "OT", "D-"]) + str(rng.randint(10000, 999999))


def gen_population(W, rng):
    goods = [v for v in W.vendors.values() if v.cat in ("PACK", "LAB", "OFF", "LOG") and v.terms != "ADV"]
    plan = [("79001", rand_day(rng, D(2024, 11, 20), D(2025, 12, 10))) for _ in range(66)]
    plan += [(rng.choice(goods).no, rand_day(rng, D(2024, 12, 2), D(2025, 12, 12))) for _ in range(168)]
    plan.sort(key=lambda x: (x[1], x[0]))
    for vno, date in plan:
        v = W.vendors[vno]
        ic = v.cat == "IC"
        nl = rng.choice([2, 3, 3, 4, 4]) if ic else rng.choice([1, 2, 2, 2, 3, 3])
        mats = rng.sample(v.mats, min(nl, len(v.mats)))
        lines = []
        for m in mats:
            mm = W.mats[m]
            q = rng.randint(mm.qmin, mm.qmax)
            if ic:
                p = rnd(mm.std * (F(103, 100) if date >= CUT and int(m) % 2 == 0 else 1))
            else:
                p = mm.std if rng.random() < 0.4 else rnd(mm.std * F(rng.randint(970, 1030), 1000))
            lines.append((m, q, p))
        lead = rng.randint(28, 45) if ic else rng.randint(6, 21)
        po = new_po(W, vno, date, lines, lead=lead, buyer=rng.choice(["E01", "E02", "E03"]))
        prom = po.lines[0].promised_true
        nship = rng.choices([1, 2, 3], [64, 28, 8])[0]
        split = {}
        for l in po.lines:
            if rng.random() < 0.6:
                split[l.pos] = split_qty(rng, l.qty, nship)
            else:
                at = 0 if rng.random() < 0.8 else nship - 1
                split[l.pos] = [l.qty if k == at else 0 for k in range(nship)]
        dd = prom + (rng.randint(-4, 0) if rng.random() < 0.80 else rng.randint(1, 14)) * DAY
        last_ship = {l.pos: max(k for k in range(nship) if split[l.pos][k]) for l in po.lines}
        for k in range(nship):
            if k:
                dd = dd + rng.randint(7, 25) * DAY
            dd = bd(dd)
            items = {p: s[k] for p, s in split.items() if s[k]}
            if not items:
                continue
            if nship > 1:
                for p in items:
                    W.k1["partial delivery / several invoices per order line"].append((po.key, p))
            deliver(W, rng, po, dd, items, {p for p in items if last_ship[p] == k})


def deliver(W, rng, po, dd, items, final):
    v = W.vendors[po.vend]
    ic = v.cat == "IC"
    gr_post = bd(dd + (0 if dd >= D(2025, 12, 24) else rng.choice([0, 0, 1, 1, 2, 3])) * DAY)
    dn = dn_no(rng, v)
    # ---- invoice timing
    if ic:
        if rng.random() < 0.85:
            inv_dt = dd - rng.randint(10, 18) * DAY
        else:
            inv_dt = dd + rng.randint(3, 20) * DAY
            for p in items:
                W.k1["group invoice after the goods"].append((po.key, p))
        inv_post = bd(inv_dt + rng.randint(2, 5) * DAY)
    else:
        inv_dt = dd + rng.randint(0, 8) * DAY
        inv_post = bd(inv_dt + rng.randint(2, 7) * DAY)
    has_gr = gr_post <= YE
    has_inv = inv_post <= YE
    gr_items = dict(items)
    price, short, clear = {}, {}, []
    feature_ok = (not ic) and po.cur == "NVK" and has_gr and has_inv and inv_post <= D(2025, 11, 20)
    for p in items:
        l = line_of(po, p)
        m = W.mats[l.mat]
        if feature_ok and p in final and rng.random() < 0.085 and items[p] >= 20:
            gr_items[p] = items[p] + max(1, items[p] // 50)
            clear.append(p)
            W.k1["free over-delivery, remainder cleared by manual journal"].append((po.key, p))
        elif feature_ok and p in final and v.base and m.conv >= 20 and rng.random() < 0.22:
            short[p] = rng.randint(1, m.conv // 5)
            clear.append(p)
            W.k1["vendor billed fewer pieces than the receipt in order units, remainder cleared"].append(
                (po.key, p))
        elif (not ic) and has_inv and rng.random() < 0.06:
            price[p] = rnd(l.price0 * F(rng.choice([94, 96, 97, 103, 104, 106, 108]), 100))
            if rng.random() < 0.5:
                l.price = price[p]
                l.chgd_on = bd(min(inv_dt, YE) - rng.randint(0, 3) * DAY)
                W.k1["price changed after the order, order amended (receipt at old price)"].append(
                    (po.key, p))
            else:
                W.k1["price changed after the order, order not amended"].append((po.key, p))
    if v.base and has_inv:
        for p in items:
            if W.mats[line_of(po, p).mat].conv > 1:
                W.k1["invoice in base unit, order and receipt in order unit"].append((po.key, p))
    # ---- goods receipt, possibly with a correction (reversal and re-posting)
    gr = None
    if has_gr:
        usr = rng.choice(["WH01", "WH01", "WH02"])
        if rng.random() < 0.085 and dd <= D(2025, 12, 12) and not clear:
            wrong = dict(gr_items)
            p0 = rng.choice(list(wrong))
            wrong[p0] = wrong[p0] + rng.choice([1, 2, 5, 10]) if rng.random() < 0.6 else max(1, wrong[p0] // 2)
            g0 = add_gr(W, po, gr_post, dd, wrong, dn, usr)
            post2 = bd(gr_post + rng.choice([0, 1, 2, 4, 6]) * DAY)
            add_rev(W, g0, post2, txt="")
            gr = add_gr(W, po, post2, dd, gr_items, dn, usr)
            for p in gr_items:
                W.k1["receipt reversed and re-posted (correction)"].append((po.key, p))
        else:
            gr = add_gr(W, po, gr_post, dd, gr_items, dn, usr)
    if ic and has_inv and not has_gr:
        for p in items:
            W.k1["group delivery in transit at year end (invoiced, not received)"].append((po.key, p))
    # ---- invoice
    if has_inv:
        inv = add_inv(W, po.vend, inv_dt, inv_post, inv_lines(W, po, items, price, short), rng)
        for p in clear:
            add_mj(W, month_end_bd(inv_post + (rng.choice([0, 0, 30]) if inv_post.month < 11 else 0) * DAY),
                   [clr(po, p)], usr="MHAL" if rng.random() < 0.12 else None)
        # return to vendor with credit note
        if (not ic) and has_gr and not clear and not price and dd <= D(2025, 10, 31) and rng.random() < 0.11:
            p0 = rng.choice(list(items))
            if items[p0] >= 4:
                rq = max(1, items[p0] * rng.randint(10, 40) // 100)
                ret = bd(max(gr.post, dd + rng.randint(10, 25) * DAY))
                idx = [i for i, it in enumerate(gr.items) if it.pos == p0][0]
                add_rev(W, gr, ret, doc=ret, qtys={idx: rq}, txt="return to vendor")
                cn_dt = ret + rng.randint(5, 15) * DAY
                add_inv(W, po.vend, cn_dt, bd(cn_dt + rng.randint(2, 6) * DAY),
                        inv_lines(W, po, {p0: rq}), rng, typ="KG")
                W.k1["return to vendor with credit note"].append((po.key, p0))
        elif (not ic) and has_gr and not clear and not price and rng.random() < 0.025:
            cn_dt = inv_dt + rng.randint(10, 30) * DAY
            if bd(cn_dt + 4 * DAY) <= YE:
                p0 = rng.choice(list(items))
                l0 = [x for x in inv.lines if x.pos == p0][0]
                add_inv(W, po.vend, cn_dt, bd(cn_dt + rng.randint(2, 4) * DAY),
                        [NS(po=po, pos=p0, mat=l0.mat, qty=0, uom=l0.uom,
                            net=rnd(l0.net * F(rng.randint(2, 5), 100)), gl="", cc="", txt="price adjustment")],
                        rng, typ="KG")
                W.k1["credit note for price only (quantity zero)"].append((po.key, p0))


def gen_services(W, rng):
    for vno, gl, cc, text, monthly, (lo, hi) in SERVICES:
        if monthly:
            dates = [D(2025, m, rng.randint(2, 9)) for m in range(1, 13)]
        else:
            dates = sorted(rand_day(rng, D(2025, 1, 6), D(2025, 12, 15)) for _ in range(rng.randint(3, 6)))
        for d in dates:
            net = rng.randint(lo, hi) * 100 + (0 if lo == hi else rng.randint(0, 99))
            ln = [NS(po=None, pos=None, mat="", qty=1, uom="", net=net, gl=gl, cc=cc,
                     txt=f"{text} {d.month:02d}/{d.year}" if monthly else text)]
            add_inv(W, vno, d, bd(d + rng.randint(1, 6) * DAY), ln, rng)


def simple_flow(W, rng, tag, vno, date, lines, gr=None, inv=None, bonus=None, lead=14, promised=None):
    """A scripted order with one receipt and one invoice. gr/inv: dates or None."""
    po = new_po(W, vno, date, lines, lead=lead, tag=tag, promised=promised)
    t = W.tags[tag]
    items = {l.pos: l.qty for l in po.lines}
    if gr:
        g = dict(items)
        for p, b in (bonus or {}).items():
            g[p] += b
        t.gr = add_gr(W, po, gr, gr, g, dn_no(rng, W.vendors[vno]))
    if inv:
        t.inv = add_inv(W, vno, inv - 4 * DAY, inv, inv_lines(W, po, items), rng)
    return po


def gen_scripted(W, rng):
    T = W.tags
    m = lambda v, i: mat_by(W, v, i)
    # --- group deliveries received in December, supply company has not invoiced yet
    for tag, date, ls, g in (("I1", D(2025, 11, 6), [(4, 60), (11, 35), (20, 48)], D(2025, 12, 12)),
                             ("I2", D(2025, 11, 13), [(7, 90), (16, 24)], D(2025, 12, 19))):
        po = new_po(W, "79001", date, [(m("79001", i), q, None) for i, q in ls], lead=(g - date).days, tag=tag,
                    buyer="E01")
        T[tag].gr = add_gr(W, po, g, g, {l.pos: l.qty for l in po.lines}, dn_no(rng, W.vendors["79001"]))
        for l in po.lines:
            W.k1["group invoice after the goods"].append((po.key, l.pos))
    # --- genuine cases of payment without receipt
    simple_flow(W, rng, "R1", "70041", D(2025, 2, 17), [(m("70041", 0), 150, 38800)], inv=D(2025, 2, 26))
    simple_flow(W, rng, "R2", "70034", D(2025, 6, 9),
                [(m("70034", 0), 20, 31500), (m("70034", 1), 30, 14200)], inv=D(2025, 6, 18))
    simple_flow(W, rng, "R3", "70043", D(2025, 8, 4), [(m("70043", 1), 15, 122000)], inv=D(2025, 8, 14))
    simple_flow(W, rng, "R4", "70045", D(2025, 11, 3), [(m("70045", 0), 80, 149000)], inv=D(2025, 11, 10))
    # --- laboratory orders, invoice paid, no receipt in the system
    simple_flow(W, rng, "K4a", "70022", D(2025, 4, 14), [(m("70022", 0), 20, 127500)], inv=D(2025, 4, 28))
    simple_flow(W, rng, "K4b", "70023", D(2025, 9, 8), [(m("70023", 0), 3, 1490000)], inv=D(2025, 9, 19))
    # --- advance payment vendors: invoice and payment first, receipt later
    adv = [("T1", "70026", D(2025, 1, 20), [(0, 12, 214000)], 21), ("T2", "70044", D(2025, 3, 3), [(0, 6, 335000)], 24),
           ("T3", "70026", D(2025, 5, 19), [(1, 40, 46800), (0, 5, 214000)], 17),
           ("T4", "70044", D(2025, 8, 25), [(1, 20, 64000)], 20), ("T5", "70026", D(2025, 10, 13), [(1, 30, 46800)], 22),
           ("T6", "70044", D(2025, 12, 1), [(0, 8, 335000)], 42)]
    for tag, vno, date, ls, wait in adv:
        po = new_po(W, vno, date, [(m(vno, i), q, p) for i, q, p in ls], lead=wait, tag=tag)
        items = {l.pos: l.qty for l in po.lines}
        T[tag].inv = add_inv(W, vno, date + 2 * DAY, bd(date + 3 * DAY), inv_lines(W, po, items), rng)
        g = bd(date + wait * DAY)
        if g <= YE:
            T[tag].gr = add_gr(W, po, g, g, items, dn_no(rng, W.vendors[vno]))
        W.k1["advance payment vendor: invoice paid before the receipt"].append((po.key, 10))
    # --- return after payment, credit note still outstanding at year end
    simple_flow(W, rng, "L1", "70013", D(2025, 10, 6), [(m("70013", 1), 200, 9600)], gr=D(2025, 10, 20),
                inv=D(2025, 10, 28))
    T["L1"].ret = add_rev(W, T["L1"].gr, D(2025, 12, 16), doc=D(2025, 12, 16), qtys={0: 50},
                          txt="return to vendor")
    # --- one invoice covering two orders
    pa = simple_flow(W, rng, "X1a", "70011", D(2025, 3, 10), [(m("70011", 0), 80, 45500)], gr=D(2025, 3, 21))
    pb = simple_flow(W, rng, "X1b", "70011", D(2025, 3, 12), [(m("70011", 2), 120, 38800)], gr=D(2025, 3, 24))
    T["X1a"].inv = add_inv(W, "70011", D(2025, 3, 26), D(2025, 3, 31),
                           inv_lines(W, pa, {10: 80}) + inv_lines(W, pb, {10: 120}), rng)
    W.k1["one invoice covering two orders"] += [(pa.key, 10), (pb.key, 10)]
    # --- duplicate invoice cases (true world: one invoice each)
    po = new_po(W, "70011", D(2025, 9, 22), [(m("70011", 3), 100, 146000)], tag="A1")
    T["A1"].gr = add_gr(W, po, D(2025, 10, 8), D(2025, 10, 8), {10: 60}, dn_no(rng, W.vendors["70011"]))
    T["A1"].inv = add_inv(W, "70011", D(2025, 10, 10), D(2025, 10, 16), inv_lines(W, po, {10: 60}), rng)
    T["A1"].gr2 = add_gr(W, po, D(2025, 12, 11), D(2025, 12, 11), {10: 40}, dn_no(rng, W.vendors["70011"]))
    simple_flow(W, rng, "A2", "70032", D(2025, 3, 10), [(m("70032", 0), 12, 20500), (m("70032", 1), 4, 53500)],
                gr=D(2025, 3, 19), inv=D(2025, 3, 27))
    simple_flow(W, rng, "A3", "70041", D(2025, 10, 20), [(m("70041", 0), 40, 38800), (m("70041", 1), 15, 46000)],
                gr=D(2025, 11, 6), inv=D(2025, 11, 14))
    # --- receipts (true world: on the right line)
    simple_flow(W, rng, "B1", "70021", D(2025, 5, 5), [(m("70021", 0), 30, 98000), (m("70021", 2), 18, 235000)],
                gr=D(2025, 5, 19), inv=D(2025, 5, 27))
    simple_flow(W, rng, "B2a", "70012", D(2025, 8, 11), [(m("70012", 2), 80, 45500)], gr=D(2025, 8, 25),
                inv=D(2025, 9, 2))
    simple_flow(W, rng, "B2b", "70012", D(2025, 9, 15), [(m("70012", 2), 80, 47200)], gr=D(2025, 9, 29),
                inv=D(2025, 10, 7))
    po = new_po(W, "70045", D(2025, 10, 27),
                [(m("70045", 0), 25, 149000), (m("70045", 1), 10, 214000), (m("70045", 2), 60, 89000)], tag="B3",
                promised=D(2025, 11, 10))
    T["B3"].gr = add_gr(W, po, D(2025, 11, 10), D(2025, 11, 10), {10: 25, 20: 10}, dn_no(rng, W.vendors["70045"]))
    T["B3"].inv = add_inv(W, "70045", D(2025, 11, 12), D(2025, 11, 17), inv_lines(W, po, {10: 25, 20: 10}), rng)
    T["B3"].gr2 = add_gr(W, po, D(2025, 12, 15), D(2025, 12, 15), {30: 60}, dn_no(rng, W.vendors["70045"]),
                         usr="WH02")
    # --- EUR order received in December, not invoiced
    simple_flow(W, rng, "C1", "70025", D(2025, 11, 17), [(m("70025", 0), 4, 745000), (m("70025", 1), 10, 66000)],
                gr=D(2025, 12, 9), lead=21)
    # --- manual clearing cases
    simple_flow(W, rng, "D1", "70014", D(2025, 7, 21), [(m("70014", 2), 60, 64000)], gr=D(2025, 8, 5))
    po = T["D1"].po
    T["D1"].inv = add_inv(W, "70014", D(2025, 8, 12), D(2025, 11, 18), inv_lines(W, po, {10: 60}), rng)
    simple_flow(W, rng, "D2", "70013", D(2025, 6, 2), [(m("70013", 0), 1000, 6150)], gr=D(2025, 6, 16),
                inv=D(2025, 6, 24), bonus={10: 40})
    add_mj(W, D(2025, 6, 30), [clr(T["D2"].po, 10)], tag="D2")
    simple_flow(W, rng, "D3a", "70015", D(2025, 8, 18), [(m("70015", 0), 120, 31800), (m("70015", 1), 40, 74250)],
                gr=D(2025, 9, 1), inv=D(2025, 9, 9), bonus={10: 1, 20: 1})
    simple_flow(W, rng, "D3b", "70023", D(2025, 8, 26), [(m("70023", 1), 20, 112640)], gr=D(2025, 9, 8),
                inv=D(2025, 9, 16), bonus={10: 1})
    add_mj(W, D(2025, 9, 30), [clr(T["D3a"].po, 10), clr(T["D3a"].po, 20), clr(T["D3b"].po, 10)], tag="D3")
    simple_flow(W, rng, "D4", "70013", D(2025, 5, 12), [(m("70013", 2), 500, 11180)], gr=D(2025, 5, 22),
                inv=D(2025, 5, 28), bonus={10: 20})
    add_mj(W, D(2025, 5, 30), [clr(T["D4"].po, 10)], tag="D4")
    simple_flow(W, rng, "D5", "70022", D(2025, 10, 6), [(m("70022", 2), 14, 138420)], gr=D(2025, 10, 17),
                inv=D(2025, 10, 24), bonus={10: 1})
    add_mj(W, D(2025, 10, 31), [clr(T["D5"].po, 10)], tag="D5")
    for tg in ("D2", "D3a", "D3b", "D4", "D5"):
        for l in T[tg].po.lines:
            W.k1["free over-delivery, remainder cleared by manual journal"].append((T[tg].po.key, l.pos))
    # --- receipts posted once (true world)
    simple_flow(W, rng, "F1", "70012", D(2025, 11, 24), [(m("70012", 0), 300, 41200), (m("70012", 1), 120, 65500)],
                gr=D(2025, 12, 10))
    simple_flow(W, rng, "F2", "70031", D(2025, 2, 3), [(m("70031", 1), 25, 26800)], gr=D(2025, 2, 12),
                inv=D(2025, 2, 20))
    # --- invoice in pieces
    simple_flow(W, rng, "H1", "70014", D(2025, 4, 7), [(m("70014", 0), 36, 89000)], gr=D(2025, 4, 22),
                inv=D(2025, 4, 29))
    # --- large receipt, vendor has not invoiced by year end (invoice arrives in January 2026)
    simple_flow(W, rng, "K3", "70043", D(2025, 8, 18), [(m("70043", 0), 220, 109000), (m("70043", 1), 60, 122000)],
                gr=D(2025, 9, 9))


def mark_k7(W, rng):
    """Promised date overwritten after a vendor announced a delay (recorded value = last agreed date)."""
    first_gr = {}
    for e in W.ev:
        if e.k == "GR" and not getattr(e, "rev_of", None):
            for it in e.items:
                first_gr.setdefault((e.po.key, it.pos), e.doc)
    for po in W.pos:
        if po.tag or po.date < CUT:
            continue
        for l in po.lines:
            fg = first_gr.get((po.key, l.pos))
            if fg and fg > l.promised_true and rng.random() < 0.34:
                l.promised_rec = bd(fg + rng.choice([0, 0, 0, 1, 2]) * DAY)
                l.chgd_on = max(l.chgd_on or po.date, bd(l.promised_true - rng.randint(0, 4) * DAY))
                W.k7.append((po.key, l.pos))


def build_true_world():
    rng = random.Random(SEED)
    W = World()
    build_master(W, rng)
    gen_population(W, rng)
    gen_services(W, rng)
    gen_scripted(W, rng)
    mark_k7(W, rng)
    schedule_payments(W)
    return W


# ---------------------------------------------------------- disorder (kind 2/3)

def dup_ref(ref, style):
    if style == 0:
        return ref.replace("-", "")
    if style == 1:
        return ref.replace("-", " ").replace("/", "-")
    return re.sub(r"-0+", "-", ref)


def apply_disorder(W):
    """Mutates W (a deep copy of the true world) into what the affiliate recorded."""
    rng = random.Random(SEED + 1)
    T = W.tags

    def dup(tag, post, style, pay):
        o = T[tag].inv
        e = add_inv(W, o.vend, o.inv_dt, post, copy.deepcopy(o.lines), rng, ref=dup_ref(o.ref, style),
                    usr="MHAL" if style == 1 else "TBER", pay=None,
                    txt="")
        assert e.ref != o.ref
        for l, lo in zip(e.lines, o.lines):
            l.po = lo.po
        if pay:
            add_pay(W, e, pay)
        T[tag].dup = e

    dup("A1", D(2025, 11, 5), 0, D(2025, 12, 4))
    dup("A2", D(2025, 4, 22), 2, D(2025, 5, 15))
    dup("A3", D(2025, 12, 12), 1, None)

    # receipts booked to the wrong order line
    g = T["B1"].gr
    it = [x for x in g.items if x.pos == 20][0]
    it.mat, it.pos = line_of(g.po, 20).mat, 10
    g = T["B2b"].gr
    g.items[0].mat = line_of(g.po, 10).mat
    g.po = T["B2a"].po
    g = T["B3"].gr2
    g.items[0].mat, g.items[0].pos = line_of(g.po, 30).mat, 10

    # currency slips
    T["C1"].po.cur = "NVK"
    T["R3"].inv.cur = "EUR"
    [p for p in W.ev if p.k == "PAY" and p.inv is T["R3"].inv][0].cur = "NVK"

    # manual journals
    T["D1"].mj = add_mj(W, D(2025, 10, 31), [clr(T["D1"].po, 10)])
    T["D2"].mj.items[0].acct = "1410"
    T["D3"].mj.lump = True
    T["D4"].mj.items[0].factor = -1
    T["D5"].mj.items[0].amt = 134820

    # mapping gaps
    W.mapping = [r for r in W.mapping if r.loc != "2815"]
    [r for r in W.mapping if r.loc == "6950" and r.vfrom == CUT][0].vfrom = D(2025, 8, 1)

    # receipts posted twice
    for tag, post in (("F1", D(2025, 12, 12)), ("F2", D(2025, 2, 13))):
        g = T[tag].gr
        T[tag].dup = add_gr(W, g.po, post, g.doc, {i.pos: i.qty for i in g.items}, g.dn, usr="WH02")

    # unit slip on an invoice
    T["H1"].inv.lines[0].uom = W.mats[T["H1"].inv.lines[0].mat].ou

    # ---- kind 3
    T["K3"].mj = add_mj(W, D(2025, 12, 19), [clr(T["K3"].po, 10), clr(T["K3"].po, 20)])
    p = [p for p in W.ev if p.k == "PAY" and p.inv is T["R2"].inv][0]
    gross = inv_total(T["R2"].inv)[2]
    s = list(f"{gross // 100}")
    s[-2], s[-3] = s[-3], s[-2]
    p.amt = int("".join(s)) * 100 + gross % 100
    assert p.amt != gross
    W.mapping.append(NS(loc="1450", name="Goods in transit", grp="120150", vfrom=CUT,
                        vto=D(9999, 12, 31), by="MHAL"))
    W.mapping.sort(key=lambda r: (r.loc, r.vfrom, r.grp))


# ------------------------------------------------------------------ the ERP

def grir_acct(W, po, post):
    return "2815" if W.vendors[po.vend].cat == "IC" and post >= CUT else "2810"


def simulate(W):
    """Play the world through the ERP model. Returns the tables."""
    T = NS(gr=[], invh=[], invl=[], pay=[], gl=[], S={}, W=W)
    mats, vendors = W.mats, W.vendors
    # ---- numbers
    old = {"PO": 45002010, "GR": 50013040, "INV": 51007300, "MJ": 10000810, "PAY": 15004400}
    new = {"PO": 0, "GR": 0, "INV": 0, "MJ": 0, "PAY": 0}
    pref = {"PO": "P25-{:04d}", "GR": "W25-{:05d}", "INV": "E25-{:05d}", "MJ": "J25-{:04d}", "PAY": "Z25-{:04d}"}
    step = {"PO": 3, "GR": 3, "INV": 2, "MJ": 9, "PAY": 2}

    def number(kind, date, salt):
        s = 1 + (salt * 7 + 3) % step[kind]
        if date < CUT:
            old[kind] += s
            return str(old[kind])
        new[kind] += s
        return pref[kind].format(new[kind])

    for po in sorted(W.pos, key=lambda p: (p.date, p.key)):
        po.no = number("PO", po.date, po.key)
        for l in po.lines:
            T.S[(po.key, l.pos)] = NS(grq=F(0), grv=0, irq=F(0), irv=0, bal=0, man=0, po=po, line=l)

    def gl(doc, typ, post, docdt, acct, amt, cur, amt_doc, ref, line, vend, txt, usr, show=True):
        if amt == 0 and amt_doc == 0:
            return
        T.gl.append(NS(doc=doc, typ=typ, post=post, docdt=docdt, acct=acct, amt=amt, cur=cur, amt_doc=amt_doc,
                       ref=ref, line=line if show else None, vend=vend, txt=txt, usr=usr))

    def settle(st):
        if st.grq > st.irq and st.grq > 0:
            target = -rnd((st.grq - st.irq) * F(st.grv) / st.grq)
        elif st.irq > st.grq and st.irq > 0:
            target = rnd((st.irq - st.grq) * F(st.irv) / st.irq)
        else:
            target = 0
        diff = target - st.bal
        st.bal = target
        return diff

    def reclass(d, reverse):
        rows = [(k, st) for k, st in sorted(T.S.items())
                if vendors[st.po.vend].cat == "IC" and st.po.date <= d and st.bal + st.man > 0]
        if not rows:
            return
        usr = "JKOV" if d <= D(2025, 7, 31) else "MHAL"
        acct = "2815" if d >= CUT else "2810"
        doc = number("MJ", d, len(T.gl))
        tot = 0
        back = []
        for k, st in rows:
            x = st.bal + st.man
            gl(doc, "SA", d, d, acct, -x, "NVK", -x, "", k, st.po.vend, "Goods in transit group", usr)
            st.man -= x
            tot += x
            back.append((k, st, x))
        gl(doc, "SA", d, d, "1450", tot, "NVK", tot, "", None, "", "Goods in transit group", usr)
        if reverse:
            d2 = d + DAY
            doc2 = number("MJ", d2, len(T.gl))
            for k, st, x in back:
                gl(doc2, "SA", d2, d2, acct, x, "NVK", x, doc, k, st.po.vend, "Rev. goods in transit group", usr)
                st.man += x
            gl(doc2, "SA", d2, d2, "1450", -tot, "NVK", -tot, doc, None, "", "Rev. goods in transit group", usr)

    order = {"GR": 1, "INV": 2, "MJ": 3, "PAY": 4}
    events = sorted(W.ev, key=lambda e: (e.post, order[e.k], e.seq))
    cuts = [(D(2024, 12, 31), True), (D(2025, 6, 30), True), (YE, False)]
    payrows = []
    for e in events + [None]:
        while cuts and (e is None or e.post > cuts[0][0]):
            reclass(*cuts.pop(0))
        if e is None:
            break
        if e.k == "GR":
            po = e.po
            e.no = number("GR", e.post, e.seq)
            rate = book(e.post) if po.cur == "EUR" else F(1)
            dacct = "6950" if po.cur == "EUR" else "5890"
            inv_tot, diff_tot = defaultdict(int), 0
            for n, it in enumerate(e.items):
                l = line_of(po, it.pos)
                st = T.S[(po.key, it.pos)]
                matno = it.mat or l.mat
                mm = mats[matno]
                if it.rev is not None:
                    lc = -rnd(F(it.rev.lc) * it.qty / it.rev.qty)
                    dc = -rnd(F(it.rev.dc) * it.qty / it.rev.qty)
                    q = -F(it.qty)
                else:
                    dc = rnd(F(it.qty) * l.price0)
                    lc = rnd(dc * rate)
                    it.lc, it.dc = lc, dc
                    q = F(it.qty)
                st.grq += q
                st.grv += lc
                st.bal -= lc
                acct = grir_acct(W, po, e.post)
                gl(e.no, "WE", e.post, e.doc, acct, -lc, po.cur, -dc, e.dn, (po.key, it.pos), po.vend,
                   mm.descr[:28], e.usr)
                inv_tot["1400" if mm.grp == "FERT" else "1410"] += lc
                diff = settle(st)
                if diff:
                    gl(e.no, "WE", e.post, e.doc, acct, diff, po.cur, 0, e.dn, (po.key, it.pos), po.vend,
                       "Diff. GR/IR", "BATCH")
                    diff_tot += diff
                T.gr.append(NS(e=e, it=it, doc=e.no, pos=n + 1, post=e.post, docdt=e.doc, mvt=it.mvt, po=po,
                               popos=it.pos, mat=matno, qty=it.qty, uom=mm.ou, qty_bu=it.qty * mm.conv,
                               bu=mm.bu, lc=abs(lc), dn=e.dn, usr=e.usr, txt=e.txt,
                               rev_ref=e.rev_of.no if getattr(e, "rev_of", None) else ""))
            for a, x in sorted(inv_tot.items()):
                gl(e.no, "WE", e.post, e.doc, a, x, po.cur, 0 if po.cur == "EUR" else x, e.dn, None, po.vend,
                   "Goods receipt " + po.no, e.usr)
            gl(e.no, "WE", e.post, e.doc, dacct, -diff_tot, po.cur, 0, e.dn, None, po.vend, "Diff. GR/IR", "BATCH")
        elif e.k == "INV":
            e.no = number("INV", e.post, e.seq)
            v = vendors[e.vend]
            sign = 1 if e.typ == "RE" else -1
            rate = book(e.post) if e.cur == "EUR" else F(1)
            net, tax, gross = inv_total(e)
            ap = "3350" if v.cat == "IC" else ("3300" if v.ctry == "NV" else "3310")
            ap_lc = rnd(gross * rate)
            gl(e.no, e.typ, e.post, e.inv_dt, ap, -sign * ap_lc, e.cur, -sign * gross, e.ref, None, e.vend,
               v.name[:24], e.usr)
            used = 0
            diffs = defaultdict(int)
            for n, l in enumerate(e.lines):
                lc = rnd(l.net * rate)
                used += lc
                if l.po is not None:
                    po = l.po
                    st = T.S[(po.key, l.pos)]
                    mm = mats[l.mat]
                    qo = F(l.qty, mm.conv) if (l.uom == mm.bu and mm.bu != mm.ou) else F(l.qty)
                    l.qo = qo * sign
                    st.irq += sign * qo
                    st.irv += sign * lc
                    st.bal += sign * lc
                    acct = grir_acct(W, po, e.post)
                    gl(e.no, e.typ, e.post, e.inv_dt, acct, sign * lc, e.cur, sign * l.net, e.ref,
                       (po.key, l.pos), e.vend, mm.descr[:28], e.usr)
                    diff = settle(st)
                    if diff:
                        gl(e.no, e.typ, e.post, e.inv_dt, acct, diff, e.cur, 0, e.ref, (po.key, l.pos), e.vend,
                           "Diff. GR/IR", "BATCH")
                        diffs["6950" if po.cur == "EUR" else "5890"] -= diff
                else:
                    gl(e.no, e.typ, e.post, e.inv_dt, l.gl, sign * lc, e.cur, sign * l.net, e.ref, None, e.vend,
                       l.txt, e.usr)
                T.invl.append(NS(e=e, l=l, doc=e.no, ln=n + 1))
            gl(e.no, e.typ, e.post, e.inv_dt, "1570", sign * (ap_lc - used), e.cur, sign * tax, e.ref, None,
               e.vend, "Input VAT", e.usr)
            for a, x in sorted(diffs.items()):
                gl(e.no, e.typ, e.post, e.inv_dt, a, x, e.cur, 0, e.ref, None, e.vend, "Diff. GR/IR", "BATCH")
            e.lc = ap_lc
            T.invh.append(NS(e=e, doc=e.no, net=net, tax=tax, gross=gross))
        elif e.k == "MJ":
            e.no = number("MJ", e.post, e.seq)
            rows = []
            for it in e.items:
                st = T.S[(it.po.key, it.pos)]
                grir = grir_acct(W, it.po, e.post)
                amt = (it.amt if it.amt is not None else -(st.bal + st.man)) * it.factor
                acct = it.acct or grir
                if acct in GRIR and not e.lump:
                    st.man += amt
                it.posted = amt
                rows.append((acct, amt, (it.po.key, it.pos), "6950" if it.po.cur == "EUR" else "5890", it.po))
            if e.lump:
                tot = sum(r[1] for r in rows)
                gl(e.no, "SA", e.post, e.post, rows[0][0], tot, "NVK", tot, "", None, "", e.txt, e.usr)
                gl(e.no, "SA", e.post, e.post, rows[0][3], -tot, "NVK", -tot, "", None, "", e.txt, e.usr)
            else:
                for acct, amt, k, cacct, po in rows:
                    gl(e.no, "SA", e.post, e.post, acct, amt, "NVK", amt, "", k, po.vend, e.txt, e.usr)
                    gl(e.no, "SA", e.post, e.post, cacct, -amt, "NVK", -amt, "", k, "", e.txt, e.usr)
        elif e.k == "PAY":
            payrows.append(e)
    # ---- payments: one document per vendor and run
    docs = {}
    for e in sorted(payrows, key=lambda e: (e.post, e.inv.vend, e.inv.post, e.inv.seq)):
        key = (e.post, e.inv.vend)
        if key not in docs:
            docs[key] = number("PAY", e.post, len(docs))
        inv = e.inv
        sign = 1 if inv.typ == "RE" else -1
        cur = e.cur or inv.cur
        amt = sign * (e.amt if e.amt is not None else inv_total(inv)[2])
        lc = rnd(amt * (book(e.post) if cur == "EUR" else 1))
        e.no = docs[key]
        iso = e.post.isocalendar()
        T.pay.append(NS(e=e, doc=e.no, run=f"{'IC' if vendors[inv.vend].cat == 'IC' else 'PR'}{iso[0]}-{iso[1]:02d}",
                        date=e.post, vend=inv.vend, inv=inv, amt=amt, cur=cur, lc=lc))
    # ---- carry forward: 2024 postings on the clearing accounts become opening items
    cf = defaultdict(int)
    keep = []
    for r in T.gl:
        if r.post < YS:
            if r.acct in GRIR:
                cf[(r.acct, r.line)] += r.amt
        else:
            keep.append(r)
    T.gl_all = T.gl
    opening = [NS(doc="09000001", typ="SV", post=YS, docdt=YS, acct=a, amt=x, cur="NVK", amt_doc=x, ref="",
                  line=k, vend=T.S[k].po.vend, txt="Balance carried forward", usr="BATCH")
               for (a, k), x in sorted(cf.items()) if x]
    T.gl = opening + keep
    return T


# ------------------------------------------------------------------ analysis

def line_balances(T):
    bal = defaultdict(int)
    for r in T.gl:
        if r.acct in GRIR:
            bal[r.line] += r.amt
    return bal


def q2_flags(T):
    """PO-based invoices paid although the invoiced quantity is not covered by receipts.
    Returns (ye, atpay): ye = {(inv seq, (po key, pos)): uncovered qty at year end},
    atpay = {(inv seq, (po key, pos))} uncovered when the payment was made."""
    W = T.W
    grs, irs = defaultdict(list), defaultdict(list)
    for r in T.gr:
        grs[(r.po.key, r.popos)].append((r.post, F(r.qty) * (1 if r.mvt == "101" else -1)))
    for x in T.invl:
        if x.l.po is not None:
            irs[(x.l.po.key, x.l.pos)].append((x.e.post, x.e.seq, x.e, x.l.qo))
    paid = {}
    for p in T.pay:
        paid[p.inv.seq] = p.date
    ye, atpay = {}, set()
    for k, lst in irs.items():
        lst.sort(key=lambda t: (t[0], t[1]))
        grtot = sum(q for _, q in grs[k])
        cn = sum(q for _, _, e, q in lst if e.typ == "KG")
        cum = cn
        for post, seq, e, q in lst:
            if e.typ != "RE":
                continue
            cum += q
            if e.seq in paid:
                un = min(q, cum - grtot)
                if un > 0:
                    ye[(e.seq, k)] = un
                gr_at = sum(g for d, g in grs[k] if d <= paid[e.seq])
                if cum - gr_at > 0:
                    atpay.add((e.seq, k))
    return ye, atpay


def q3_measure(T, field, datefield, h2_only=True):
    """Value received on lines with a promised date, and the part delivered after it."""
    reversed_qty = defaultdict(F)
    for r in T.gr:
        if r.it.rev is not None:
            reversed_qty[id(r.it.rev)] += F(r.qty)
    tot = late = tot_e = late_e = 0
    for r in T.gr:
        if r.it.rev is not None:
            continue
        if h2_only and r.po.date < CUT:
            continue
        prom = getattr(line_of(r.po, r.popos), field)
        if prom is None:
            continue
        net = F(r.qty) - reversed_qty[id(r.it)]
        if net <= 0:
            continue
        val = F(r.lc) * net / F(r.qty)
        eur = val / book(r.post)
        tot += val
        tot_e += eur
        if getattr(r, datefield) > prom:
            late += val
            late_e += eur
    return NS(tot=rnd(tot), late=rnd(late), tot_e=rnd(tot_e), late_e=rnd(late_e),
              share=float(late / tot) if tot else 0.0)


def q3_overdue(T, field):
    """Open quantity at year end on lines whose promised date has passed (value at order price)."""
    out = 0
    for (k, pos), st in T.S.items():
        prom = getattr(st.line, field)
        if prom is None or st.po.date < CUT or prom > YE:
            continue
        open_q = F(st.line.qty) - st.grq
        if open_q > 0:
            v = open_q * st.line.price0
            out += v * (book(YE) if st.po.cur == "EUR" else 1)
    return rnd(out)


# ------------------------------------------------------------------- output

def dmy(d):
    return d.strftime("%d.%m.%Y")


def iso(d):
    return d.isoformat()


def assign_str(T, line, post):
    if line is None:
        return ""
    po = T.S[line].po
    return f"{po.no}{line[1]:05d}" if post < CUT else f"{po.no}/{line[1]}"


def write_csv(path, header, rows):
    buf = io.StringIO(newline="")
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(header)
    w.writerows(rows)
    path.write_bytes(buf.getvalue().encode("utf-8"))
    return len(rows)


def write_data(T, out):
    """Write the delivered extracts. Ordinary noise is added here and nowhere recorded."""
    W = T.W
    nz = random.Random(SEED + 2)
    out.mkdir(parents=True, exist_ok=True)
    counts = {}

    # vendors
    rows = []
    for v in W.vendors.values():
        city = v.city.upper() if nz.random() < 0.12 else v.city
        email = "" if nz.random() < 0.2 else \
            f"{nz.choice(['office', 'invoices', 'sales', 'info'])}@{v.name.split()[0].lower()}.example.com"
        vat = f"NV{nz.randint(10000000, 99999999)}" if v.ctry == "NV" else ""
        bank = f"{nz.randint(100, 999)}-{nz.randint(1000000, 9999999)}-{nz.randint(10, 99)}" if v.ctry == "NV" else ""
        rows.append([v.no, v.name, city, v.ctry, v.cur, v.terms, "MSC" if v.cat == "IC" else "", vat, bank, email,
                     dmy(D(2016 + int(v.no) % 8, 1 + int(v.no) % 12, 1 + int(v.no) % 27)), ""])
    counts["vendors.csv"] = write_csv(out / "vendors.csv",
                                      ["vend_no", "name1", "city", "ctry", "cur", "pay_trm", "grp_partner", "vat_no",
                                       "bank_acct", "email", "crt_dt", "blk"], rows)
    # materials
    rows = []
    for m in W.mats.values():
        d = m.descr
        r = nz.random()
        if r < 0.05:
            d = d.upper()
        elif r < 0.09:
            d = d.replace(" ", "  ", 1)
        rows.append([m.no, d, m.grp, m.bu, m.ou, m.conv, money(m.std), m.cur,
                     "" if nz.random() < 0.3 else nz.choice(["A", "B", "C"])])
    counts["materials.csv"] = write_csv(out / "materials.csv",
                                        ["mat_no", "descr", "mat_grp", "base_uom", "ord_uom", "conv_fact",
                                         "std_price", "price_cur", "abc"], rows)
    # purchase orders
    hrows, lrows = [], []
    for po in sorted(W.pos, key=lambda p: (p.date, p.key)):
        sts = [T.S[(po.key, l.pos)] for l in po.lines]
        done = all(s.grq >= s.line.qty and s.irq >= s.grq for s in sts)
        v = W.vendors[po.vend]
        hrows.append([po.no, dmy(po.date), po.vend, po.cur, "" if nz.random() < 0.08 else po.buyer, v.terms,
                      po.crt, "C" if done else "O"])
        for l in po.lines:
            st = T.S[(po.key, l.pos)]
            m = W.mats[l.mat]
            lrows.append([po.no, l.pos, l.mat, m.descr[:30], l.qty, m.ou, money(l.price), 1,
                          money(rnd(F(l.qty) * l.price)), dmy(l.promised_rec) if l.promised_rec else "",
                          dmy(l.chgd_on) if l.chgd_on else "", "X" if st.grq >= l.qty else "",
                          "V2" if v.ctry == "NV" else "V0"])
    counts["po_header.csv"] = write_csv(out / "po_header.csv",
                                        ["po_no", "po_dt", "vend_no", "cur", "buyer", "pay_trm", "crt_by", "stat"],
                                        hrows)
    counts["po_lines.csv"] = write_csv(out / "po_lines.csv",
                                       ["po_no", "pos", "mat_no", "short_txt", "qty", "ou", "price", "per",
                                        "net_val", "dlv_dt", "chgd_on", "dlv_compl", "tax_cd"], lrows)
    # goods receipts
    rows = []
    notes = ["ok", "2 pal.", "carton dented, accepted", "after 15h", "ramp 2", "checked w/ purchasing",
             "driver waited", "cold chain ok", "partial", "see mail"]
    for r in sorted(T.gr, key=lambda r: (r.post, r.e.seq, r.pos)):
        txt = r.txt or (nz.choice(notes) if nz.random() < 0.07 else "")
        dn = r.dn
        x = nz.random()
        if x < 0.02:
            dn = ""
        elif x < 0.035:
            dn = "w/o DN"
        elif x < 0.06:
            dn = dn.lower()
        uom = r.uom.lower() if nz.random() < 0.02 else r.uom
        rows.append([r.doc, r.pos, dmy(r.post), dmy(r.docdt), r.mvt, r.po.no, r.popos, r.mat, qfmt(r.qty), uom,
                     qfmt(r.qty_bu), r.bu, money(r.lc), dn, r.rev_ref, r.usr, txt])
    counts["goods_receipts.csv"] = write_csv(out / "goods_receipts.csv",
                                             ["gr_doc", "gr_pos", "post_dt", "doc_dt", "mvt", "po_no", "po_pos",
                                              "mat_no", "qty", "uom", "qty_bu", "bu", "amt_lc", "dn_no", "rev_ref",
                                              "usr", "txt"], rows)
    # invoices
    hrows, lrows = [], []
    htxt = ["scan 2", "orig. by post", "urgent", "via requester", "e-invoice", "pdf by mail"]
    for h in sorted(T.invh, key=lambda h: (h.e.post, h.e.seq)):
        e = h.e
        ref = e.ref
        x = nz.random()
        if x < 0.02:
            ref = ref + " "
        elif x < 0.035:
            ref = " " + ref
        hrows.append([h.doc, e.typ, e.vend, ref, iso(e.inv_dt), iso(e.post), e.cur, money(h.gross), money(h.tax),
                      iso(e.due), "", e.usr, e.txt or (nz.choice(htxt) if nz.random() < 0.09 else "")])
    for x in sorted(T.invl, key=lambda x: (x.e.post, x.e.seq, x.ln)):
        l, e = x.l, x.e
        cc = l.cc if nz.random() > 0.1 else ""
        lrows.append([x.doc, x.ln, l.po.no if l.po is not None else "", l.pos if l.po is not None else "", l.mat,
                      qfmt(l.qty), l.uom, money(l.net), "V2" if e.taxrate else "V0", l.gl, cc, l.txt])
    counts["inv_header.csv"] = write_csv(out / "inv_header.csv",
                                         ["inv_doc", "doc_type", "vend_no", "vend_ref", "inv_dt", "post_dt", "cur",
                                          "gross", "tax", "due_dt", "pay_blk", "usr", "hdr_txt"], hrows)
    counts["inv_lines.csv"] = write_csv(out / "inv_lines.csv",
                                        ["inv_doc", "ln", "po_no", "po_pos", "mat_no", "qty", "uom", "net", "tax_cd",
                                         "gl_acct", "cc", "txt"], lrows)
    # payments
    rows = []
    for p in sorted(T.pay, key=lambda p: (p.date, p.doc, p.inv.post, p.inv.seq)):
        name = W.vendors[p.vend].name
        x = nz.random()
        if x < 0.15:
            name = name.upper()
        elif x < 0.27:
            name = name.replace(" d.o.o.", "").replace(" d.d.", "")
        elif x < 0.33:
            name = name[:18]
        elif x < 0.38:
            name = name.replace("d.o.o.", "doo")
        rows.append([p.run, p.doc, dmy(p.date), p.vend, name, p.inv.no, money(p.amt), p.cur, money(p.lc),
                     "NETT" if p.run.startswith("IC") else "TRF", "1110" if p.cur == "EUR" else "1100"])
    counts["payments.csv"] = write_csv(out / "payments.csv",
                                       ["run_id", "pay_doc", "pay_dt", "vend_no", "payee", "inv_doc", "amt", "cur",
                                        "amt_lc", "meth", "bank_gl"], rows)
    # ledger lines
    rows = []
    ln = defaultdict(int)
    mjtxt = ["GR/IR clearing", "GR/IR clearing", "clearing GR/IR diff", "GR/IR diff cleared", "Clear. GR/IR"]
    for r in T.gl:
        if r.acct not in EXTRACT_ACCTS:
            continue
        ln[r.doc] += 1
        txt = r.txt
        if r.typ == "SA" and txt == "GR/IR clearing":
            txt = mjtxt[(int(r.doc[-3:]) + len(r.doc)) % len(mjtxt)]
        elif r.typ in ("RE", "KG", "WE") and nz.random() < 0.06:
            txt = ""
        rows.append([r.doc, ln[r.doc], iso(r.post), iso(r.docdt), f"{r.post.year}/{r.post.month:02d}", r.typ, r.acct,
                     "D" if r.amt > 0 or (r.amt == 0 and r.amt_doc > 0) else "C", money(abs(r.amt)), r.cur,
                     money(abs(r.amt_doc)), r.ref, assign_str(T, r.line, r.post), r.vend, txt, r.usr])
    counts["gl_lines.csv"] = write_csv(out / "gl_lines.csv",
                                       ["doc_no", "ln", "post_dt", "doc_dt", "per", "doc_type", "acct", "dc",
                                        "amt_lc", "cur", "amt_doc", "ref", "assign", "vend_no", "txt", "usr"], rows)
    # exchange rates
    rows = [[f"2025-{m:02d}", "EUR", "NVK", 1, f"{float(book(D(2025, m, 15))):.4f}", EOM[(2025, m)]]
            for m in range(1, 13)]
    counts["fx_rates.csv"] = write_csv(out / "fx_rates.csv",
                                       ["per", "from_cur", "to_cur", "units", "rate_book", "rate_eom"], rows)
    counts["acct_mapping.xlsx"] = write_mapping(W, out / "acct_mapping.xlsx")
    write_pdfs(out)
    return counts


def write_mapping(W, path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mapping"
    ws.append(["loc_acct", "loc_descr", "grp_acct", "grp_descr", "valid_from", "valid_to", "chg_by"])
    n = 0
    for r in sorted(W.mapping, key=lambda r: (r.loc, r.vfrom, r.grp)):
        ws.append([r.loc, r.name, r.grp, GROUP_ACCTS[r.grp], dt.datetime.combine(r.vfrom, dt.time()),
                   dt.datetime.combine(r.vto, dt.time()), r.by])
        n += 1
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=6):
        for c in row:
            c.number_format = "DD.MM.YYYY"
    ws2 = wb.create_sheet("GroupAccounts")
    ws2.append(["grp_acct", "grp_descr"])
    for k, v in sorted(GROUP_ACCTS.items()):
        ws2.append([k, v])
    stamp = dt.datetime(2025, 6, 27, 16, 40, 0)
    wb.properties.created = stamp
    wb.properties.modified = stamp
    wb.properties.creator = "JKOV"
    wb.properties.lastModifiedBy = "JKOV"
    buf = io.BytesIO()
    wb.save(buf)
    # repack with fixed timestamps so the file is byte-identical on every run
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in src.namelist():
            zi = zipfile.ZipInfo(name, date_time=(2025, 6, 27, 16, 40, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            data = src.read(name)
            if name == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(created|modified)[^>]*>)[^<]*", rb"\g<1>2025-06-27T16:40:00Z", data)
            dst.writestr(zi, data)
    return n


def write_pdfs(out):
    import textwrap
    import fitz
    import docs_text
    docs = [("posting_instruction_local.pdf", docs_text.POSTING_INSTRUCTION,
             "LI-FIN-04 Purchasing, goods receipt and invoice verification", "D:20230301090000+01'00'"),
            ("group_accounting_manual_ch7.pdf", docs_text.GROUP_MANUAL,
             "Group Accounting Manual - Chapter 7 (extract)", "D:20250115090000+01'00'"),
            ("handover_note_accounting.pdf", docs_text.HANDOVER_NOTE,
             "Handover note purchasing ledger", "D:20250728153000+02'00'")]
    for fname, blocks, title, stamp in docs:
        doc = fitz.open()
        page, y = None, 0
        for style, text in blocks:
            size, font, width, gap = {"h1": (14, "hebo", 70, 6), "h2": (11, "hebo", 88, 4),
                                      "meta": (9, "heit", 100, 3), "p": (10, "helv", 92, 7)}[style]
            lines = textwrap.wrap(text, width)
            need = len(lines) * (size + 3.5) + gap + (10 if style == "h2" else 0)
            if page is None or y + need > 790:
                page = doc.new_page(width=595, height=842)
                y = 60
            if style == "h2":
                y += 10
            for ln in lines:
                y += size + 3.5
                page.insert_text((62, y), ln, fontsize=size, fontname=font)
            y += gap
        doc.set_metadata({"title": title, "author": "Meridia Pharma Norvania d.o.o.", "producer": "",
                          "creator": "", "creationDate": stamp, "modDate": stamp})
        doc.save(out / fname, garbage=0, deflate=False, no_new_id=True)
        doc.close()


SOURCES = [
    ("vendors", "csv", "vendors.csv", "Vendor master of the affiliate, including the group supply company."),
    ("materials", "csv", "materials.csv", "Material master with base unit, order unit and the conversion between them."),
    ("po_header", "csv", "po_header.csv", "Purchase order headers: vendor, order date, currency."),
    ("po_lines", "csv", "po_lines.csv", "Purchase order lines: material, quantity in order unit, price, delivery date."),
    ("goods_receipts", "csv", "goods_receipts.csv", "Goods receipt lines against purchase order lines, including reversals."),
    ("inv_header", "csv", "inv_header.csv", "Vendor invoice and credit note headers, with and without purchase order."),
    ("inv_lines", "csv", "inv_lines.csv", "Vendor invoice and credit note lines."),
    ("payments", "csv", "payments.csv", "Payment runs and the invoices and credit notes they settled."),
    ("gl_lines", "csv", "gl_lines.csv", "Ledger lines on the inventory, clearing, vendor and difference accounts, automatic and manual."),
    ("acct_mapping", "xlsx", "acct_mapping.xlsx", "Mapping of local accounts to group accounts, with validity dates."),
    ("fx_rates", "csv", "fx_rates.csv", "Monthly exchange rates NVK per EUR for FY2025: booking rate and month-end rate."),
    ("posting_instruction_local", "pdf", "posting_instruction_local.pdf", "The affiliate's local posting instruction for purchasing, goods receipt and invoice verification."),
    ("group_accounting_manual_ch7", "pdf", "group_accounting_manual_ch7.pdf", "Extract of the group accounting manual: goods receipt, invoice verification, clearing account, manual journals."),
    ("handover_note_accounting", "pdf", "handover_note_accounting.pdf", "Handover note of the accountant who left the affiliate at the end of July 2025."),
]
QUESTIONS = [
    "What is the balance of goods received but not yet invoiced at year end, in EUR?",
    "Which vendor invoices were paid without a goods receipt, and for how much?",
    "How much of the purchase volume was delivered late?",
]


def write_manifest(out):
    lines = ["""# The affiliate landscape — a small country affiliate with its own local ERP.
#
# Meridia Pharma Norvania d.o.o.: a fictional sales and distribution company of
# the fictional Meridia Pharma Group, in the fictional Republic of Norvania
# (currency NVK). Nine extracts from its local ERP as they arrive in the group
# data warehouse, and three documents. Fiscal year 2025; on 1 July 2025 the
# affiliate changed its document numbering and its account mapping and did not
# restate the earlier months.
#
# Why it exists: realistic in the way that matters — not everything fits, and
# nobody wrote down why.

name: affiliate
domain: finance
frozen: true
seed: 20251003

# This landscape is read with a foundation document rather than a domain guide;
# the packaged finance pack is named here only because the manifest requires one.
domain_guide:
  packaged: finance

# DO NOT OPEN. The answer key states which cases were seeded and what the right
# answers are; the generator contains the same in code. Reading either one ends
# the blind evaluation permanently.
answer_key: answer-key/ANSWER_KEY.md
generator: generator/generate_affiliate.py

questions:"""]
    for q in QUESTIONS:
        lines.append(f'  - "{q}"')
    lines += ["", "sources:"]
    for name, kind, path, descr in SOURCES:
        sha = hashlib.sha256((out / path).read_bytes()).hexdigest()
        lines += [f"  - name: {name}", f"    kind: {kind}", f"    path: {path}", f"    sha256: {sha}",
                  f'    description: "{descr}"']
    (ROOT / "manifest.yaml").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    import selfcheck
    import answer_key
    true_w = build_true_world()
    rec_w = copy.deepcopy(true_w)
    clean = simulate(true_w)
    results = selfcheck.run_checks(true_w, clean)
    bad = [r for r in results if not r[1]]
    if bad:
        for name, ok, detail in bad:
            print("SELFCHECK FAILED:", name, detail)
        sys.exit(1)
    apply_disorder(rec_w)
    rec = simulate(rec_w)
    out = ROOT / "data"
    for p in out.glob("*"):
        p.unlink()
    counts = write_data(rec, out)
    answer_key.write(true_w, clean, rec_w, rec, ROOT / "answer-key" / "ANSWER_KEY.md")
    write_manifest(out)
    print(f"selfcheck on clean base: {len(results)} checks passed")
    for k, v in counts.items():
        print(f"  {k:22s} {v:5d} rows")
    print(f"  total rows {sum(counts.values())}, files {len(list(out.glob('*')))}")


if __name__ == "__main__":
    main()
